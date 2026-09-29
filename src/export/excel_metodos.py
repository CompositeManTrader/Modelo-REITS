"""Los métodos de valuación trimestre por trimestre, en Excel.

Una hoja por emisor con cada fin de trimestre: los insumos en azul sobre amarillo
(precio, flujo, dividendo, tasas) y todo lo que se deriva de ellos —múltiplo, yields,
primas— como fórmula viva, con la convención de ``export.excel``. Los percentiles van
como valor: salen de la historia MENSUAL completa y reconstruirlos con fórmulas sobre los
renglones trimestrales daría otro número. Los umbrales de barato y caro están en la hoja
Léeme y mueven todas las señales.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.estudio.metodos import BARATO, CARO, METODOS, NOTAS_DE_METODO, ResultadoMetodos, trimestral
from src.export.excel import (
    FMT_BPS,
    FMT_MONEDA,
    FMT_PCT,
    FMT_PCT1,
    FMT_VECES,
    FUENTE_FORMULA,
    FUENTE_SECCION,
    FUENTE_TITULO,
    RELLENO_SECCION,
    escribir_etiqueta,
    escribir_input,
)

LEEME = "Léeme"
UMBRAL_BARATO = f"'{LEEME}'!$B$6"
UMBRAL_CARO = f"'{LEEME}'!$B$7"

# (encabezado, columna del panel o None si es fórmula, formato)
INSUMOS = (
    ("Precio", "precio", FMT_MONEDA),
    ("Flujo por acción, 12 meses", "flujo", FMT_MONEDA),
    ("Dividendo anualizado", "dividendo", FMT_MONEDA),
    ("Treasury 10 años", "treasury_10a", FMT_PCT),
    ("Inflación 12 meses", "inflacion_12m", FMT_PCT),
    ("Bonos Baa", "baa", FMT_PCT),
)
# Columnas: A trimestre, B fecha, C–H insumos, I–P fórmulas.
FORMULAS = (
    ("Múltiplo precio/flujo", "=IFERROR(C{f}/D{f},\"\")", FMT_VECES),
    ("Yield de flujo", "=IFERROR(D{f}/C{f},\"\")", FMT_PCT),
    ("Yield de dividendo", "=IFERROR(E{f}/C{f},\"\")", FMT_PCT),
    ("Tasa real", "=IFERROR(F{f}-G{f},\"\")", FMT_PCT),
    ("Prima sobre el Treasury", "=IFERROR(J{f}-F{f},\"\")", FMT_PCT),
    ("Prima sobre la tasa real", "=IFERROR(J{f}-L{f},\"\")", FMT_PCT),
    ("Prima sobre bonos Baa", "=IFERROR(J{f}-H{f},\"\")", FMT_PCT),
    ("Dividendo sobre el Treasury", "=IFERROR(K{f}-F{f},\"\")", FMT_PCT),
)


def _encabezados(ws: Worksheet, textos: list[str], fila: int = 1) -> None:
    for j, texto in enumerate(textos, start=1):
        c = ws.cell(row=fila, column=j, value=texto)
        c.font = FUENTE_SECCION
        c.fill = RELLENO_SECCION


def _valor(ws: Worksheet, fila: int, col: int, valor, formato: str | None = None) -> None:
    c = ws.cell(row=fila, column=col, value=None if valor is None or pd.isna(valor) else valor)
    c.font = FUENTE_FORMULA
    if formato:
        c.number_format = formato


def _senal(celda: str) -> str:
    return f'=IF({celda}="","",IF({celda}>={UMBRAL_BARATO},"barato",IF({celda}<{UMBRAL_CARO},"caro","medio")))'


def _hoja_leeme(wb: Workbook, r: ResultadoMetodos) -> None:
    ws = wb.active
    ws.title = LEEME
    ws["A1"] = "Métodos de valuación en el tiempo"
    ws["A1"].font = FUENTE_TITULO
    ws["A2"] = f"Emisores: {', '.join(r.paneles)} · datos al {r.hasta:%d-%m-%Y} · generado el {dt.date.today():%d-%m-%Y}"
    ws["A3"] = ("Azul sobre amarillo: insumo, editable. Negro: fórmula o valor calculado por el estudio. "
                "Herramienta de análisis, no asesoría de inversión.")
    escribir_etiqueta(ws, "A5", "Umbrales de la señal (percentil contra la historia propia)", seccion=True)
    escribir_etiqueta(ws, "A6", "Barato desde el percentil")
    escribir_input(ws, "B6", BARATO, FMT_PCT1)
    escribir_etiqueta(ws, "A7", "Caro abajo del percentil")
    escribir_input(ws, "B7", CARO, FMT_PCT1)
    escribir_etiqueta(ws, "A9", "Los métodos", seccion=True)
    for i, m in enumerate(METODOS, start=10):
        ws.cell(row=i, column=1, value=m.nombre).font = FUENTE_SECCION
        ws.cell(row=i, column=2, value=m.descripcion)
    fila = 10 + len(METODOS) + 1
    escribir_etiqueta(ws, f"A{fila}", "Cómo se calculó", seccion=True)
    for i, nota in enumerate(NOTAS_DE_METODO, start=fila + 1):
        ws.cell(row=i, column=1, value="•")
        ws.cell(row=i, column=2, value=nota)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 140


def _hoja_emisor(wb: Workbook, r: ResultadoMetodos, ticker: str) -> None:
    ws = wb.create_sheet(ticker[:31])
    q = trimestral(r.paneles[ticker]).dropna(subset=["precio"])
    encabezados = (["Trimestre", "Fecha"] + [h for h, _, _ in INSUMOS] + [h for h, _, _ in FORMULAS]
                   + [f"Percentil: {m.nombre}" for m in METODOS]
                   + ["Señal del consenso", "Retorno anual, 1 año después", "Retorno anual, 5 años después"])
    _encabezados(ws, encabezados)
    col_consenso = 2 + len(INSUMOS) + len(FORMULAS) + len(METODOS)
    for i, (f, fila) in enumerate(q.iterrows(), start=2):
        _valor(ws, i, 1, f"{f.year}-T{(f.month - 1) // 3 + 1}")
        _valor(ws, i, 2, f.date(), "dd-mm-yyyy")
        for j, (_, clave, formato) in enumerate(INSUMOS, start=3):
            v = fila[clave]
            if pd.notna(v):
                escribir_input(ws, f"{get_column_letter(j)}{i}", float(v), formato)
        for j, (_, formula, formato) in enumerate(FORMULAS, start=3 + len(INSUMOS)):
            c = ws.cell(row=i, column=j, value=formula.format(f=i))
            c.font, c.number_format = FUENTE_FORMULA, formato
        for j, m in enumerate(METODOS, start=3 + len(INSUMOS) + len(FORMULAS)):
            _valor(ws, i, j, fila[f"p_{m.clave}"], "0%")
        c = ws.cell(row=i, column=col_consenso + 1, value=_senal(f"{get_column_letter(col_consenso)}{i}"))
        c.font = FUENTE_FORMULA
        _valor(ws, i, col_consenso + 2, fila["adelante_1a"], FMT_PCT1)
        _valor(ws, i, col_consenso + 3, fila["adelante_5a"], FMT_PCT1)
    ws.freeze_panes = "C2"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 11
    for j in range(3, len(encabezados) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14


def _hoja_hoy(wb: Workbook, r: ResultadoMetodos) -> None:
    ws = wb.create_sheet("Hoy")
    tickers = list(r.paneles)
    _encabezados(ws, ["Método", *[f"Percentil {t}" for t in tickers], *[f"Señal {t}" for t in tickers]])
    h = r.hoy()
    for i, m in enumerate(METODOS, start=2):
        ws.cell(row=i, column=1, value=m.nombre).font = FUENTE_FORMULA
        for j, t in enumerate(tickers, start=2):
            v = h[(h["emisor"] == t) & (h["clave"] == m.clave)]["percentil"].iloc[0]
            _valor(ws, i, j, v, "0%")
            c = ws.cell(row=i, column=j + len(tickers), value=_senal(f"{get_column_letter(j)}{i}"))
            c.font = FUENTE_FORMULA
    ws.column_dimensions["A"].width = 36


def _hoja_asignacion(wb: Workbook, r: ResultadoMetodos) -> None:
    ws = wb.create_sheet("Asignación")
    tickers = list(r.paneles)
    a = r.asignacion()
    ws["A1"] = "A cuál de los tres va la aportación del mes: al de percentil más alto contra su propia historia"
    ws["A1"].font = FUENTE_TITULO
    cab = ["Método", "Desde", "TIR", "TIR partes iguales", "Ventaja", "Ventaja al más caro",
           "Ventaja 1ª mitad", "Ventaja 2ª mitad", "Azar que la iguala", "Cambios de emisor", "Hoy va a"]
    _encabezados(ws, cab, fila=3)
    for i, m in enumerate(METODOS, start=4):
        x = a.loc[m.clave]
        _valor(ws, i, 1, m.nombre)
        _valor(ws, i, 2, x["desde"].date(), "mm-yyyy")
        _valor(ws, i, 3, x["tir_usd"], FMT_PCT)
        _valor(ws, i, 4, x["tir_partes_iguales"], FMT_PCT)
        for j, clave in enumerate(("ventaja", "ventaja_el_mas_caro", "ventaja_primera_mitad",
                                   "ventaja_segunda_mitad"), start=5):
            # En puntos base ya escalados: el formato bps pega el texto, no multiplica.
            _valor(ws, i, j, round(float(x[clave]) * 1e4), FMT_BPS)
        _valor(ws, i, 9, x["azar_que_le_gana"], FMT_PCT1)
        _valor(ws, i, 10, int(x["cambios"]))
        _valor(ws, i, 11, x["eleccion_hoy"])
    inicio = 4 + len(METODOS) + 2
    ws.cell(row=inicio - 1, column=1, value="Cada fin de trimestre, con el consenso de los siete").font = FUENTE_SECCION
    _encabezados(ws, ["Trimestre", *tickers, "La aportación va a"], fila=inicio)
    pcs = pd.DataFrame({t: trimestral(p)["p_consenso"] for t, p in r.paneles.items()}).dropna(how="all")
    k = len(tickers)
    ultima = get_column_letter(1 + k)
    for i, (f, fila) in enumerate(pcs.iterrows(), start=inicio + 1):
        _valor(ws, i, 1, f"{f.year}-T{(f.month - 1) // 3 + 1}")
        for j, t in enumerate(tickers, start=2):
            _valor(ws, i, j, fila[t], "0%")
        rango = f"B{i}:{ultima}{i}"
        c = ws.cell(row=i, column=2 + k, value=(
            f'=IF(COUNT({rango})<{k},"—",INDEX($B${inicio}:${ultima}${inicio},MATCH(MAX({rango}),{rango},0)))'))
        c.font = FUENTE_FORMULA
    ws.column_dimensions["A"].width = 36
    for j in range(2, 12):
        ws.column_dimensions[get_column_letter(j)].width = 15


def _hoja_evaluacion(wb: Workbook, r: ResultadoMetodos) -> None:
    ws = wb.create_sheet("Evaluación")
    ws["A1"] = "¿Predicen? Percentil contra el retorno de los años siguientes, un renglón por trimestre"
    ws["A1"].font = FUENTE_TITULO
    cab = ["Método", "Correlación 1 año", "Correlación 5 años", "5 años si barato", "5 años si medio",
           "5 años si caro", "Brecha barato − caro", "Caro contra el efectivo", "Veces que caro le ganó al efectivo",
           "TIR con el método (solo ese papel)", "TIR sin reglas", "Ventaja"]
    fila = 3
    for t in r.paneles:
        ws.cell(row=fila, column=1, value=f"{r.nombres[t]} ({t})").font = FUENTE_SECCION
        _encabezados(ws, cab, fila=fila + 1)
        e, b = r.evaluaciones[t], r.backtests[t]
        for i, m in enumerate(METODOS, start=fila + 2):
            x, y = e.loc[m.clave], b.loc[m.clave]
            _valor(ws, i, 1, m.nombre)
            _valor(ws, i, 2, round(float(x["rho_1a"]), 3), "0.00")
            _valor(ws, i, 3, round(float(x["rho_5a"]), 3), "0.00")
            for j, c in enumerate(("r5_barato", "r5_medio", "r5_caro", "r5_barato_menos_caro",
                                   "caro_contra_efectivo", "caro_le_gana_al_efectivo"), start=4):
                _valor(ws, i, j, x[c], FMT_PCT1)
            _valor(ws, i, 10, y["tir_usd"], FMT_PCT)
            _valor(ws, i, 11, y["tir_sin_reglas"], FMT_PCT)
            _valor(ws, i, 12, round(float(y["ventaja_tir"]) * 1e4), FMT_BPS)
        fila += len(METODOS) + 4
    ws.column_dimensions["A"].width = 36
    for j in range(2, 13):
        ws.column_dimensions[get_column_letter(j)].width = 15


def exportar(r: ResultadoMetodos, ruta: Path | str) -> Path:
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _hoja_leeme(wb, r)
    _hoja_hoy(wb, r)
    _hoja_asignacion(wb, r)
    _hoja_evaluacion(wb, r)
    for t in r.paneles:
        _hoja_emisor(wb, r, t)
    wb.save(ruta)
    return ruta
