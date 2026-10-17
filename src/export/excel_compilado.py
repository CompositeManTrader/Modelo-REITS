"""El Excel del compilado: las cifras de cada análisis y una calculadora del plan.

Una hoja por tema con las tablas que citan los PDFs del compilado, y la hoja «Mi plan» con
fórmulas vivas: los insumos van en azul sobre amarillo (convención de ``export.excel``) y todo
lo demás se recalcula al cambiarlos. Las tablas de resultados van como valor: salen de las
corridas guardadas y reconstruirlas con fórmulas no tendría sentido.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

from src.export.excel import (
    FMT_MILES,
    FMT_MONEDA,
    FMT_PCT,
    FMT_PCT1,
    FUENTE_FORMULA,
    FUENTE_SECCION,
    FUENTE_TITULO,
    RELLENO_SECCION,
    escribir_etiqueta,
    escribir_formula,
    escribir_input,
)

FMT_PB = '+#,##0;"−"#,##0;0'
FMT_ANIOS = '0.0" años"'
FMT_MESES = '0" meses"'
# Yield de 12 meses de SCHH en el estudio de ETFs (data/estudios/etfs/resultados/indice_actual.csv).
YIELD_SCHH = 0.031

# Insumos de «Mi plan»: (etiqueta, valor de ejemplo, formato, nota). Fila 5 en adelante.
INSUMOS_PLAN = (
    ("Ahorro mensual (pesos)", 4_000, FMT_MILES, "Lo que te queda cada mes después de gastos. Tú dijiste: menos de 5,000."),
    ("Gastos mensuales (pesos)", 15_000, FMT_MILES, "Lo que gastas en un mes normal (ejemplo)."),
    ("Meses de gastos del fondo de emergencia", 3, "0", "Entre 3 y 6; 3 es el mínimo razonable."),
    ("Fondo de emergencia que ya tienes (pesos)", 0, FMT_MILES, "Lo que ya tienes en CETES o similar."),
    ("A inversión mientras juntas el fondo (pesos)", 1_000, FMT_MILES, "Opcional, para mantener el hábito. 0 = todo al fondo."),
    ("Parte del dinero invertido que va a REITs (SCHH)", 0.20, FMT_PCT1, "Entre 10% y 25% es defendible; es criterio, no resultado."),
    ("Tipo de cambio (pesos por dólar)", 18.50, FMT_MONEDA, "Ejemplo: pon el del día que compres."),
    ("Precio de SCHH (dólares)", None, FMT_MONEDA, "Último cierre de los datos versionados; actualízalo."),
    ("Rendimiento anual esperado en dólares", 0.07, FMT_PCT1, "Supuesto para proyectar; la historia de REITs y acciones ronda 7% a 10%."),
    ("Tu tasa marginal de ISR", 0.35, FMT_PCT1, "No la sabías: se usaron 30% y 35%."),
    ("Umbral del impuesto sucesorio de EE. UU. (dólares)", 60_000, FMT_MILES, "Activos de EE. UU. arriba de esto pagan hasta 40% al heredarse."),
)
FILA_INSUMOS = 5
CELDA = {clave: f"$B${FILA_INSUMOS + i}" for i, clave in enumerate(
    ("ahorro", "gastos", "meses", "fondo", "mientras", "reits", "tc", "precio", "rend", "marginal", "umbral"))}
TASA_MENSUAL = f"((1+{CELDA['rend']})^(1/12)-1)"

# (etiqueta, fórmula, formato, explicación). Van debajo de los insumos.
RESULTADOS_PLAN = (
    ("Meta del fondo de emergencia (pesos)", f"={CELDA['gastos']}*{CELDA['meses']}", FMT_MILES, "Gastos × meses."),
    ("Lo que falta del fondo (pesos)", f"=MAX(0,{CELDA['gastos']}*{CELDA['meses']}-{CELDA['fondo']})", FMT_MILES, ""),
    ("Meses para completar el fondo", "=IF({falta}=0,0,IF({ahorro}-{mientras}<=0,\"nunca: baja lo que va a inversión\","
     "ROUNDUP({falta}/({ahorro}-{mientras}),0)))", FMT_MESES, "Con el ahorro menos lo que mandas a inversión mientras."),
    ("Etapa", "=IF({falta}=0,\"Invierte todo el ahorro\",\"Primero el fondo de emergencia\")", "@", ""),
    ("Con el fondo completo: a SCHH cada mes (pesos)", f"={CELDA['ahorro']}*{CELDA['reits']}", FMT_MILES, "Ahorro × parte a REITs."),
    ("Con el fondo completo: a acciones amplias (pesos)", f"={CELDA['ahorro']}*(1-{CELDA['reits']})", FMT_MILES, "El resto."),
    ("SCHH cada mes (dólares)", f"={CELDA['ahorro']}*{CELDA['reits']}/{CELDA['tc']}", FMT_MONEDA, "Pesos ÷ tipo de cambio."),
    ("Acciones de SCHH cada mes", f"={CELDA['ahorro']}*{CELDA['reits']}/{CELDA['tc']}/{CELDA['precio']}", "0.0",
     "Si es menos de 1 y tu broker no vende fracciones, junta 2 o 3 meses."),
    ("Tasa del dividendo con broker de EE. UU.", f"={CELDA['marginal']}+0.09", FMT_PCT1,
     "Marginal + 9 puntos (10% adicional menos el crédito; modelo del repo, por confirmar)."),
    ("Años para que todo lo invertido llegue al umbral sucesorio", f"=NPER({TASA_MENSUAL},-{CELDA['ahorro']}/{CELDA['tc']},0,"
     f"{CELDA['umbral']})/12", FMT_ANIOS, "Si todo el ahorro va a activos de EE. UU. (SCHH y acciones de EE. UU.)."),
    ("Años para que solo SCHH llegue al umbral", f"=NPER({TASA_MENSUAL},-{CELDA['ahorro']}*{CELDA['reits']}/{CELDA['tc']},0,"
     f"{CELDA['umbral']})/12", FMT_ANIOS, ""),
)


def _titulo(ws: Worksheet, titulo: str, bajada: str) -> None:
    ws["A1"] = titulo
    ws["A1"].font = FUENTE_TITULO
    ws["A2"] = bajada
    ws["A2"].alignment = Alignment(wrap_text=False)


def _tabla(ws: Worksheet, fila: int, encabezado: list[str], filas: list[list], formatos: list[str | None],
           *, seccion: str = "") -> int:
    """Escribe una tabla de valores desde ``fila``; devuelve la fila siguiente libre."""
    if seccion:
        escribir_etiqueta(ws, f"A{fila}", seccion, seccion=True)
        fila += 1
    for j, h in enumerate(encabezado, start=1):
        c = ws.cell(row=fila, column=j, value=h)
        c.font = FUENTE_SECCION
        c.fill = RELLENO_SECCION
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, f in enumerate(filas, start=fila + 1):
        for j, (v, fmt) in enumerate(zip(f, formatos, strict=True), start=1):
            c = ws.cell(row=i, column=j, value=None if v is None or (isinstance(v, float) and pd.isna(v)) else v)
            c.font = FUENTE_FORMULA
            c.alignment = Alignment(wrap_text=isinstance(v, str) and len(v) > 40, vertical="top")
            if fmt:
                c.number_format = fmt
    return fila + len(filas) + 2


def _anchos(ws: Worksheet, anchos: list[int]) -> None:
    for j, a in enumerate(anchos):
        ws.column_dimensions[chr(ord("A") + j)].width = a


def _pb(v) -> float | None:
    return None if v is None or pd.isna(v) else round(float(v) * 1e4)


def _preguntas(wb: Workbook) -> None:
    from src.export.compilado import DOCUMENTOS, PREGUNTAS

    ws = wb.active
    ws.title = "Preguntas"
    _titulo(ws, "Compilado del análisis de REITs",
            f"Cada pregunta que se hizo, su respuesta corta y qué documento la explica. Generado el "
            f"{dt.date.today():%d-%m-%Y}. Herramienta de análisis, no asesoría de inversión ni fiscal.")
    fila = _tabla(ws, 4, ["Tema", "La pregunta", "Respuesta corta", "Documento"],
                  [[t, q, r, d] for t, q, r, d in PREGUNTAS], [None] * 4)
    _tabla(ws, fila, ["#", "Documento", "Archivo", "Qué responde"], [list(d) for d in DOCUMENTOS], [None] * 4,
           seccion="Los documentos del compilado (en docs/compilado/)")
    _anchos(ws, [14, 60, 70, 34])
    ws.freeze_panes = "A5"


def _conclusiones(wb: Workbook) -> None:
    from src.export.compilado import CONCLUSIONES

    ws = wb.create_sheet("Conclusiones")
    _titulo(ws, "Las conclusiones", "Lo que se encontró, en orden. Favorable = la evidencia apoya usarlo.")
    tono = {"favorable": "Favorable", "desfavorable": "Desfavorable", "neutral": "Neutral"}
    _tabla(ws, 4, ["Conclusión", "Por qué", "Tono"], [[t, x, tono[s]] for t, x, s in CONCLUSIONES], [None] * 3)
    _anchos(ws, [44, 110, 14])


def _cuando(wb: Workbook, res: dict) -> None:
    ws = wb.create_sheet("Cuándo entrar")
    _titulo(ws, "¿Cuándo entrar? La investigación de timing",
            "Desarrollo 1972-2015, FTSE Nareit All Equity, perfil SIC del modelo. Mejora = TIR de la regla menos la "
            "de aportar siempre, en pb al año. Criterio: +50 pb robusto.")
    d = res["fase5"]["desarrollo"].sort_values("mejora_rebalanceo", ascending=False)
    fila = _tabla(ws, 4, ["Señal", "Familia", "Regla", "Mejora (pb)", "Ejecutando un mes tarde (pb)",
                          "Con doble comisión (pb)", "Solo el dinero nuevo (pb)", "Caída máxima", "Pasa"],
                  [[r.senal, r.familia, r.regla, _pb(r.mejora_rebalanceo), _pb(r.mejora_rebalanceo_con_rezago),
                    _pb(r.mejora_rebalanceo_doble_costo), _pb(r.mejora_nuevo), r.caida_rebalanceo,
                    "Sí" if r.pasa else "No"] for r in d.itertuples()],
                  [None, None, None, FMT_PB, FMT_PB, FMT_PB, FMT_PB, FMT_PCT1, None],
                  seccion="Las 22 reglas de entrada (fase 5)")
    t = res["fase3"]["frecuencia"]
    fila = _tabla(ws, fila, ["Horizonte (meses)", "Ventanas independientes", "El efectivo ganó",
                             "REITs (mediana anual)", "Efectivo (mediana anual)"],
                  [[int(r.horizonte_meses), int(r.independientes), r.fraccion_efectivo_gana, r.rendimiento_reits_mediana,
                    r.rendimiento_efectivo_mediana] for r in t.itertuples()],
                  ["0", "0", FMT_PCT1, FMT_PCT1, FMT_PCT1], seccion="Qué tan seguido gana esperar en efectivo (fase 3)")
    t = res["fase3"]["techo"]
    fila = _tabla(ws, fila, ["Regla con información perfecta", "Modo", "TIR", "Contra aportar siempre (pb)",
                             "Caída máxima"],
                  [[r.regla, r.modo, r.tir, r.contra_aportar_siempre_bps, r.caida_maxima] for r in t.itertuples()],
                  [None, None, FMT_PCT, FMT_PB, FMT_PCT1], seccion="El techo: qué ganaría un oráculo (fase 3)")
    _tabla(ws, fila, ["Pregunta", "Mínimo detectable (pb al año)", "Qué se vio"],
           [["Cuándo entrar", "195 a 432", "Una ventaja de 100 pb pasaría inadvertida 45% a 82% de las veces"],
            ["En cuáles REITs", "59 a 110", "Momentum se distingue del azar, pero queda corto"],
            ["Escoger entre O, NNN y WPC", "54", "+72 pb observado; más que 500 elecciones al azar"],
            ["Salirse a ciegas", "cuesta 86 a 302", "Lo que una regla tiene que ganar solo para empatar"]],
           [None] * 3, seccion="Análisis de poder (docs/investigacion/poder_resultados.md)")
    _anchos(ws, [34, 16, 40, 14, 18, 16, 16, 14, 8])


def _cuales(wb: Workbook, res: dict) -> None:
    from src.investigacion import fase6

    ws = wb.create_sheet("En cuáles")
    _titulo(ws, "¿En cuáles REITs? Selección entre todos los REITs de EE. UU.",
            "Aportando cada trimestre sin vender, contra repartir entre todos por partes iguales. Mejora en pb al año.")
    fila = 4
    for muestra, nombre in (("desarrollo", "Desarrollo 2011-2015"), ("validacion", "Validación 2016-2026"),
                            ("final", "Prueba final: emisores sellados")):
        d = res["fase6"][muestra]
        if muestra == "desarrollo":
            d = d.sort_values("mejora", ascending=False)
        fila = _tabla(ws, fila, ["Regla", "Escogidos (promedio)", "TIR", "TIR de todos", "Mejora (pb)",
                                 "Un trimestre tarde (pb)", "Recortan dividendo", "Recortan (todos)", "Pasa"],
                      [[fase6.NOMBRES.get(r.regla, r.regla), r.escogidos_promedio, r.tir, r.tir_todos, _pb(r.mejora),
                        _pb(r.mejora_con_rezago), r.recortes, r.recortes_todos,
                        "Sí" if getattr(r, "pasa", r.mejora >= 0.005) else "No"]
                       for r in d.itertuples()],
                      [None, "0", FMT_PCT, FMT_PCT, FMT_PB, FMT_PB, FMT_PCT1, FMT_PCT1, None], seccion=nombre)
    _anchos(ws, [40, 14, 10, 12, 12, 16, 14, 14, 8])


def _retorno(wb: Workbook) -> None:
    from src.export.compilado import EFECTIVO_5A, RETORNO_ESPERADO, descomposicion

    ws = wb.create_sheet("Retorno esperado")
    _titulo(ws, "Cuánto esperar ganar según qué tan caro está",
            "O, NNN y WPC juntos; retorno anual de los 5 años siguientes, neto con broker de EE. UU. al 35%. Ventanas "
            "encimadas y emisores que sobrevivieron: guía, no promesa.")
    fila = _tabla(ws, 4, ["Tramo", "Percentil del consenso", "Mediana anual", "Malo (10%)", "Bueno (90%)"],
                  [list(r) for r in RETORNO_ESPERADO] + [["Efectivo (T-bills)", "—", EFECTIVO_5A, None, None]],
                  [None, None, FMT_PCT1, FMT_PCT1, FMT_PCT1])
    partes = [descomposicion("O"), descomposicion("NNN")]
    _tabla(ws, fila, ["Emisor", "Desde", "Hasta", "Retorno anual", "Dividendo", "Crecimiento del dividendo",
                      "Revaluación", "Yield al inicio", "Yield hoy"],
           [[d["ticker"], d["desde"].strftime("%Y-%m"), d["hasta"].strftime("%Y-%m"), d["retorno"], d["dividendo"],
             d["crecimiento"], d["revaluacion"], d["yield_inicio"], d["yield_hoy"]] for d in partes],
           [None, None, None] + [FMT_PCT1] * 6,
           seccion="De dónde salió el rendimiento: (1 + retorno) = (1 + dividendo)(1 + crecimiento)(1 + revaluación)")
    _anchos(ws, [22, 22, 14, 14, 14, 16, 14, 14, 12])


def _etfs(wb: Workbook) -> None:
    from src.estudio import etfs

    r = etfs.cargar()
    ws = wb.create_sheet("ETFs")
    _titulo(ws, "USRT, SCHH y VNQ contra el índice de la investigación (FTSE Nareit All Equity)",
            "Retorno total mensual (Yahoo). Cada ETF desde que sigue su índice actual. Diferencia en pb al año.")
    a = r.tablas["indice_actual"]
    fila = _tabla(ws, 4, ["Serie", "Desde", "Rendimiento anual", "Contra el índice (pb)", "Se desvía al año",
                          "Caída máxima", "Peor año", "Yield 12m", "Pagos al año", "Costo"],
                  [[f.serie, f.desde, f.rendimiento_anual, _pb(f.diferencia_anual), f.error_de_seguimiento,
                    f.caida_maxima, f.peor_anio, f.yield_12m, f.pagos_por_anio, f.costo_publicado]
                   for f in a.itertuples()],
                  [None, None, FMT_PCT, FMT_PB, FMT_PCT, FMT_PCT1, FMT_PCT1, FMT_PCT, "0", FMT_PCT],
                  seccion="Con su índice actual")
    n = r.tablas["anios"]
    cols = [c for c in ("Índice", "SCHH", "USRT", "VNQ") if c in n]
    _tabla(ws, fila, ["Año", *cols, "Meses"], [[int(f["anio"]), *[f[c] for c in cols], int(f["meses"])]
                                               for _, f in n.iterrows()],
           ["0"] + [FMT_PCT1] * len(cols) + ["0"], seccion="Año por año (retorno total)")
    _anchos(ws, [26, 12, 14, 16, 14, 12, 10, 10, 10, 8])


def _satelite(wb: Workbook) -> None:
    from src.export.compilado import _satelite as cargar

    s = cargar()
    ws = wb.create_sheet("Satélite")
    _titulo(ws, "Núcleo y satélite contra SCHH (exploratoria, 2011-2026)",
            "Aportando 3,000 dólares cada trimestre sin vender. Mejora = TIR neta del satélite menos la de la "
            "referencia; cartera = 85% referencia y 15% satélite. Datos ya vistos: no confirma nada.")
    nombres = {"indice": "Índice que sigue SCHH hoy", "SCHH": "SCHH real"}
    _tabla(ws, 4, ["Satélite", "Contra", "Escogidos", "TIR del satélite", "TIR de la referencia", "Mejora (pb)",
                   "Cartera 85/15 (pb)", "Años que ganó", "De", "Caída del satélite"],
           [[f.satelite, nombres.get(f.referencia, f.referencia), f.escogidos_promedio, f.tir_satelite, f.tir_referencia,
             _pb(f.mejora), _pb(f.mejora_cartera), int(f.anios_ganados), int(f.anios), f.caida_satelite]
            for f in s.itertuples()],
           [None, None, "0", FMT_PCT, FMT_PCT, FMT_PB, FMT_PB, "0", "0", FMT_PCT1])
    fila = 6 + len(s)
    _tabla(ws, fila, ["Qué", "Cuándo"],
           [["Regla congelada", "Momentum (meses 12 a 3), 15% de cada aportación"],
            ["Aportaciones en papel", "Al cierre de feb, may, ago y nov, de nov-2026 a ago-2029 (12)"],
            ["Veredicto", "30-nov-2029: éxito con +50 pb de TIR neta y exceso bruto positivo"],
            ["Abandono", "Si después de 8 aportaciones va 200 pb abajo"]],
           [None, None], seccion="La prueba en papel (docs/estudios/satelite_preregistro.md)")
    _anchos(ws, [36, 24, 10, 14, 16, 12, 16, 12, 6, 14])


def _edges(wb: Workbook) -> None:
    from src.export.compilado import EDGES, FAMILIAS

    ws = wb.create_sheet("Edges")
    _titulo(ws, "La búsqueda de estrategias contra SCHH",
            "285 combinaciones consideradas, 14 ideas escritas con mecanismo y fuentes, 5 sobrevivieron al adversario, "
            "3 en el catálogo. Rango propio en pb al año contra SCHH.")
    _tabla(ws, 4, ["Familia", "Idea", "Rango propio (pb)", "Veredicto", "Puntaje (de 5)"],
           [[FAMILIAS[f], n, r, v, s] for f, n, r, v, s in EDGES], [None, None, None, None, "0.00"])
    _anchos(ws, [16, 56, 18, 30, 14])


def _impuestos(wb: Workbook) -> None:
    from src.export.compilado import PERFILES, PREGUNTAS_CONTADOR

    ws = wb.create_sheet("Impuestos")
    _titulo(ws, "Impuestos y canal", "Perfiles fiscales del modelo (src/investigacion/diseno.py). Por confirmar con "
            "contador.")
    fila = _tabla(ws, 4, ["Canal", "Dividendo", "Intereses", "Ganancia", "Comisión"], [list(p) for p in PERFILES],
                  [None] * 5)
    fila = _tabla(ws, fila, ["Qué", "Con broker de EE. UU.", "Lectura"],
                  [["Aportar siempre", "−145 a −184 pb al año", "Costo del canal, no de una estrategia"],
                   ["Tendencia de 10 meses (8 mercados)", "−49 a −54 pb", "Sigue rechazada"],
                   ["Momentum en los sellados", "+41 / +42 pb", "Sigue debajo de +50"],
                   ["Consenso entre O, NNN y WPC", "+60 / +63 pb", "Sigue inconcluso"]],
                  [None] * 3, seccion="Qué cambió con el broker (docs/investigacion/sensibilidad_broker.md)")
    fila = _tabla(ws, fila, ["Vehículo", "Sin vender nunca", "Vendiendo todo al final", "Si aplicara REFIPRE"],
                  [["IUSI (iShares, físico; IUSIN en la BMV)", "+21 a +25 pb", "+6 a +11 pb", "−58 pb"],
                   ["XRES (Invesco, sintético)", "+57 a +59 pb", "−38 a −53 pb", "—"]],
                  [None] * 4, seccion="ETF irlandés de acumulación contra SCHH por broker, después de impuestos")
    _tabla(ws, fila, ["#", "Pregunta para el contador"], [[i, q] for i, q in enumerate(PREGUNTAS_CONTADOR, start=1)],
           ["0", None], seccion="Las cinco preguntas")
    _anchos(ws, [50, 22, 24, 22, 12])


def _plan(wb: Workbook, precio_schh: float) -> None:
    ws = wb.create_sheet("Mi plan")
    _titulo(ws, "Mi plan de aportación",
            "Cambia solo las celdas azules sobre fondo amarillo: todo lo demás se recalcula. Los valores que trae son "
            "un ejemplo con lo que nos dijiste (menos de 5,000 pesos de ahorro y sin fondo de emergencia).")
    escribir_etiqueta(ws, f"A{FILA_INSUMOS - 1}", "Tus números", seccion=True)
    for i, (etiqueta, valor, fmt, nota) in enumerate(INSUMOS_PLAN, start=FILA_INSUMOS):
        escribir_etiqueta(ws, f"A{i}", etiqueta)
        escribir_input(ws, f"B{i}", precio_schh if valor is None else valor, fmt)
        ws[f"C{i}"] = nota
    inicio = FILA_INSUMOS + len(INSUMOS_PLAN) + 1
    escribir_etiqueta(ws, f"A{inicio}", "Lo que sale", seccion=True)
    falta = f"$B${inicio + 2}"
    ref = {"falta": falta, "ahorro": CELDA["ahorro"], "mientras": CELDA["mientras"]}
    for i, (etiqueta, formula, fmt, nota) in enumerate(RESULTADOS_PLAN, start=inicio + 1):
        escribir_etiqueta(ws, f"A{i}", etiqueta)
        escribir_formula(ws, f"B{i}", formula.format(**ref), fmt)
        ws[f"C{i}"] = nota
    fila = inicio + len(RESULTADOS_PLAN) + 2
    escribir_etiqueta(ws, f"A{fila}", "Proyección (con el fondo completo, aportando cada mes y reinvirtiendo)",
                      seccion=True)
    for j, h in enumerate(["Años", "Aportado a SCHH (dólares)", "Valor de SCHH (dólares)", "Todo lo invertido (dólares)",
                           "Dividendo de SCHH al año (dólares)"], start=1):
        c = ws.cell(row=fila + 1, column=j, value=h)
        c.font = FUENTE_SECCION
        c.fill = RELLENO_SECCION
    usd_schh = f"{CELDA['ahorro']}*{CELDA['reits']}/{CELDA['tc']}"
    for k, anios in enumerate((1, 2, 3, 5, 10, 15, 20, 25, 30), start=fila + 2):
        ws.cell(row=k, column=1, value=anios).font = FUENTE_FORMULA
        escribir_formula(ws, f"B{k}", f"={usd_schh}*12*A{k}", FMT_MILES)
        escribir_formula(ws, f"C{k}", f"=FV({TASA_MENSUAL},12*A{k},-{usd_schh},0)", FMT_MILES)
        escribir_formula(ws, f"D{k}", f"=FV({TASA_MENSUAL},12*A{k},-{CELDA['ahorro']}/{CELDA['tc']},0)", FMT_MILES)
        escribir_formula(ws, f"E{k}", f"=C{k}*{YIELD_SCHH}", FMT_MILES)
    nota = fila + 2 + 9 + 1
    ws[f"A{nota}"] = (f"El dividendo usa el yield de 12 meses de SCHH ({YIELD_SCHH:.1%}) antes de impuestos. La "
                      "proyección supone el rendimiento de arriba cada año; en la realidad hubo años de −68%.")
    ws[f"A{nota + 1}"] = "Herramienta de análisis, no asesoría de inversión ni fiscal."
    _anchos(ws, [56, 22, 26, 26, 30])
    ws.column_dimensions["C"].width = 70


def _glosario(wb: Workbook) -> None:
    from src.export.compilado import GLOSARIO

    ws = wb.create_sheet("Glosario")
    _titulo(ws, "Glosario", "Los términos que aparecen en el compilado.")
    _tabla(ws, 4, ["Término", "Qué es"], [list(g) for g in GLOSARIO], [None, None])
    _anchos(ws, [34, 110])


def precio_schh() -> float:
    """Último cierre de SCHH en los datos versionados del estudio de ETFs."""
    ruta = Path(__file__).resolve().parents[2] / "data" / "estudios" / "etfs" / "precios_mensuales.csv.gz"
    p = pd.read_csv(ruta)
    return round(float(p[p["ticker"] == "SCHH"].sort_values("fecha")["cierre"].iloc[-1]), 2)


def exportar(ruta: Path | str, res: dict) -> Path:
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _preguntas(wb)
    _conclusiones(wb)
    _plan(wb, precio_schh())
    _cuando(wb, res)
    _cuales(wb, res)
    _retorno(wb)
    _etfs(wb)
    _satelite(wb)
    _edges(wb)
    _impuestos(wb)
    _glosario(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.save(ruta)
    return ruta
