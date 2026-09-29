"""Métodos de valuación en el tiempo (``src/estudio/metodos.py``).

* La asignación entre emisores, sobre escenarios sintéticos cuyo resultado se sabe:
  repartir en tercios es la suma de tres benchmarks; mandarlo todo a uno es el
  benchmark de ese uno; el azar conserva las rachas.
* Los paneles, sobre la historia versionada: el percentil de cada método no mira al
  futuro (P1), el CPI se conoce cuando se publica, y el Excel reproduce con sus
  fórmulas lo que calculó el estudio.
"""

from __future__ import annotations

import datetime as dt
import shutil
import subprocess
import sys
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest

RAIZ = Path(__file__).resolve().parent.parent
for ruta in (str(RAIZ), str(RAIZ / "app")):
    if ruta not in sys.path:
        sys.path.insert(0, ruta)

from src.estudio import macro, mercado, metodos, reglas  # noqa: E402
from src.estudio.reglas import Parametros, Variante  # noqa: E402
from src.portafolio.metricas import tir  # noqa: E402

D = dt.date
SIN_FRICCION = Parametros(impuesto_dividendo=0.0, impuesto_ganancia=0.0, impuesto_intereses=0.0, comision=0.0)


# --------------------------------------------------------------------------------------
# Asignación, sobre escenarios sintéticos
# --------------------------------------------------------------------------------------


def _papel(precios: list[float], dividendos=()):
    """Un papel que cotiza el primer y el penúltimo día de cada mes desde 2020."""
    meses = pd.period_range("2020-01", periods=len(precios) // 2 + 1, freq="M")
    sesiones = []
    for m in meses:
        sesiones += [m.start_time.normalize(), (m.end_time - pd.Timedelta(days=1)).normalize()]
    sesiones = pd.DatetimeIndex(sesiones[: len(precios)])
    tabla = pd.DataFrame({"precio_base": precios, "usdmxn": 20.0, "rt_usd_neto": np.nan}, index=sesiones)
    div = pd.DataFrame([{"fecha_ex": pd.Timestamp(f), "monto_base": m, "monto_pagado": m, "tipo": "regular"}
                        for f, m in dividendos], columns=["fecha_ex", "monto_base", "monto_pagado", "tipo"])
    return SimpleNamespace(ticker="ZZ", tabla=tabla,
                           historia_mercado=SimpleNamespace(dividendos=div, escisiones=()))


@pytest.fixture
def tres():
    n = 16
    return {
        "AAA": _papel([10 + i for i in range(n)], dividendos=[("2020-03-01", 0.5), ("2020-06-01", 0.5)]),
        "BBB": _papel([20 - 0.5 * i for i in range(n)], dividendos=[("2020-04-01", 1.0)]),
        "CCC": _papel([15 + (i % 3) for i in range(n)]),
    }


def _sin_senal(e) -> reglas.Senales:
    fines = reglas.fines_de_mes(e.tabla.index)
    return reglas.Senales("ZZ", pd.DataFrame({"decision": "SIN SEÑAL", "disparadores": ""}, index=fines),
                          pd.DataFrame())


def test_la_tir_rapida_es_la_del_proyecto():
    rng = np.random.default_rng(3)
    fechas = pd.DatetimeIndex(sorted(pd.Timestamp("2001-01-01") + pd.to_timedelta(rng.integers(0, 9000, 60), "D")))
    montos = np.append(-rng.uniform(100, 1_000, 59), 90_000.0)
    esperado = tir(pd.Series(montos, index=fechas))
    assert metodos.tir_rapida(fechas, montos) == pytest.approx(esperado, abs=1e-9)


def test_partes_iguales_es_la_suma_de_tres_benchmarks(tres):
    """Cada emisor recibe un tercio y reinvierte sus propios dividendos: es aportar un
    tercio a cada uno por separado, al centavo."""
    m = metodos.mercado_comun(tres)
    meses = np.arange(len(m.fines))
    p = replace(SIN_FRICCION, aportacion=900.0)
    c = metodos.correr(m, meses, np.full(len(meses), -1), propio=True, parametros=p)
    tercio = replace(SIN_FRICCION, aportacion=300.0)
    suma = sum(reglas.simular(e, _sin_senal(e), pd.Series(dtype=float), variante=Variante.BENCHMARK,
                              parametros=tercio).valor_final_neto for e in tres.values())
    assert c.valor_neto == pytest.approx(suma, rel=1e-12)
    assert c.aportado == 900.0 * len(meses)


def test_todo_a_uno_es_el_benchmark_de_ese_uno(tres):
    """Si solo ese paga dividendos, mandarle todo equivale a aportarle solo a él."""
    solo = {"AAA": tres["AAA"], "CCC": tres["CCC"]}
    m = metodos.mercado_comun(solo)
    meses = np.arange(len(m.fines))
    c = metodos.correr(m, meses, np.zeros(len(meses), dtype=int), parametros=PARAMS_REALES)
    b = reglas.simular(solo["AAA"], _sin_senal(solo["AAA"]), pd.Series(dtype=float),
                       variante=Variante.BENCHMARK, parametros=PARAMS_REALES)
    assert c.valor_neto == pytest.approx(b.valor_final_neto, rel=1e-12)
    assert c.tir_usd == pytest.approx(b.tir_usd, abs=1e-9)


PARAMS_REALES = Parametros()


def test_el_azar_conserva_las_rachas():
    rng = np.random.default_rng(0)
    rachas = np.array([5, 1, 7, 3, 2, 9])
    for _ in range(50):
        x = metodos._al_azar(rachas, 3, rng)
        assert len(x) == rachas.sum()
        cortes = np.flatnonzero(np.diff(x) != 0) + 1
        assert list(np.diff(np.concatenate([[0], cortes, [len(x)]]))) == list(rachas)


def test_la_asignacion_manda_todo_al_percentil_mas_alto(tres):
    m = metodos.mercado_comun(tres)
    n = len(m.fines)
    # AAA barato los primeros meses, BBB después; CCC siempre en medio.
    paneles = {
        "AAA": pd.DataFrame({"p_multiplo": np.r_[np.full(n // 2, 0.9), np.full(n - n // 2, 0.1)]}, index=m.fines),
        "BBB": pd.DataFrame({"p_multiplo": np.r_[np.full(n // 2, 0.1), np.full(n - n // 2, 0.9)]}, index=m.fines),
        "CCC": pd.DataFrame({"p_multiplo": 0.5}, index=m.fines),
    }
    pc = metodos.percentiles_en(m, paneles, "multiplo")
    assert list(np.argmax(pc, axis=1)) == [0] * (n // 2) + [1] * (n - n // 2)
    with pytest.raises(ValueError, match="dos años"):
        metodos.asignar(m, paneles, "multiplo", n_azar=2)


# --------------------------------------------------------------------------------------
# Paneles sobre la historia versionada
# --------------------------------------------------------------------------------------


def test_el_cpi_se_conoce_cuando_se_publica():
    """El CPI de enero (fechado el 1-ene) se publica a mediados de febrero."""
    serie = pd.Series([100.0, 101.0], index=pd.DatetimeIndex(["2020-01-01", "2020-02-01"]))
    antes, despues = macro.conocido_en("cpi", serie, pd.DatetimeIndex(["2020-02-14", "2020-02-15"]))
    assert np.isnan(antes) and despues == 100.0


def _armar(repo, ticker: str, asof: dt.date):
    if not mercado.hay_estudio(ticker):
        pytest.skip(f"No está versionado el estudio de {ticker}.")
    from src.estudio import estudio

    return estudio.armar(repo, ticker, asof=asof)


def test_el_percentil_de_cada_metodo_no_mira_al_futuro(repo_sembrado):
    tarde = metodos.panel(_armar(repo_sembrado, "O", D(2026, 6, 30)))
    temprano = metodos.panel(_armar(repo_sembrado, "O", D(2015, 6, 30)))
    columnas = [f"p_{m.clave}" for m in metodos.METODOS] + ["yield_flujo", "inflacion_12m", "baa"]
    pd.testing.assert_frame_equal(tarde.loc[:"2015-05-31", columnas], temprano.loc[:"2015-05-31", columnas])


def test_el_consenso_es_el_promedio_de_los_siete(repo_sembrado):
    p = metodos.panel(_armar(repo_sembrado, "O", D(2026, 6, 30)))
    siete = p[[f"p_{c}" for c in metodos.SIETE]]
    con_cuatro = siete.notna().sum(axis=1) >= 4
    pd.testing.assert_series_equal(p.loc[con_cuatro, "p_consenso"], siete[con_cuatro].mean(axis=1),
                                   check_names=False)
    assert p.loc[~con_cuatro, "p_consenso"].isna().all()
    assert set(p["s_consenso"]) <= {"barato", "medio", "caro", "sin dato"}


@pytest.fixture(scope="module")
def r_metodos(tmp_path_factory):
    from src.datos.repositorio import Repositorio
    from src.datos.semilla import sembrar

    repo = Repositorio(ruta=tmp_path_factory.mktemp("base") / "sembrada.db")
    sembrar(repo, tickers=["O", "PLD", "GNL"], inicio=D(2018, 1, 1), fin=D(2026, 6, 30))
    tickers = [t for t in ("O", "NNN") if mercado.hay_estudio(t)]
    if len(tickers) < 2:
        pytest.skip("Hacen falta dos emisores con estudio versionado.")
    return metodos.estudiar({t: _armar(repo, t, D(2026, 6, 30)) for t in tickers}, n_azar=10)


def test_las_conclusiones_estan_completas(r_metodos):
    import re

    c = metodos.conclusiones(r_metodos)
    assert [x.titulo for x in c] == [
        "¿Los métodos de valuación predicen?", "Entonces, ¿por qué esperar a que esté barato no paga?",
        "Donde sí sirve: a cuál de los dos va la aportación", "Qué hacer hoy", "Qué tanto se le puede creer"]
    for x in c:
        # Una cifra faltante se escribe «—» o «nan» donde iba un número.
        assert not re.search(r"\bnan\b|—(?: pb|%| al año)|de — a", x.texto), x.texto


def test_la_conclusion_no_presume_controles_que_no_pasaron(r_metodos):
    """Con dos emisores y el corte de 2026-06 el azar iguala al consenso a veces: el texto
    no puede decir que los controles lo respaldan."""
    a = r_metodos.asignacion().loc["consenso"]
    texto = metodos.conclusiones(r_metodos)[2].texto
    if a["azar_que_le_gana"] > 0.05 or a["ventaja_el_mas_caro"] >= 0:
        assert "no alcanzan para respaldarlo" in texto
    else:
        assert "Tres controles lo respaldan" in texto


def test_la_asignacion_trae_sus_controles(r_metodos):
    t = r_metodos.asignacion()
    assert list(t.index) == [m.clave for m in metodos.METODOS]
    for a in r_metodos.asignaciones:
        assert len(a.azar) == 10
        assert [n.split(" (")[0] for n, _ in a.mitades] == ["primera mitad", "segunda mitad"]
        assert a.eleccion_hoy in r_metodos.paneles
        assert sum(a.meses_por_emisor.values()) > 24


def test_las_vistas_ponen_cada_fraccion_en_porcentaje(r_metodos):
    from comun import familia_de_columna

    from src.estudio import vistas_metodos as vm

    tablas = {"juntos": vm.juntos(r_metodos), "asignacion": vm.asignacion(r_metodos), "metodos": vm.metodos()}
    for nombre, construir in vm.VISTAS.items():
        tablas[nombre] = construir(r_metodos, next(iter(r_metodos.paneles)))
    problemas = []
    for nombre, tabla in tablas.items():
        if list(tabla.columns) != list(vm.COLUMNAS[nombre]):
            problemas.append(f"{nombre}: columnas {list(tabla.columns)}")
        for columna in tabla.columns:
            serie = pd.to_numeric(tabla[columna], errors="coerce").dropna()
            if serie.empty or tabla[columna].dtype == object or pd.api.types.is_integer_dtype(tabla[columna]):
                continue
            if columna.startswith("correlacion"):
                assert familia_de_columna(columna, serie) == "numero", columna
                continue
            if serie.abs().max() < 1 and familia_de_columna(columna, serie) != "porcentaje":
                problemas.append(f"{nombre}.{columna}: fracción en «{familia_de_columna(columna, serie)}»")
    assert not problemas, "\n".join(problemas)
    assert len(vm.hoy(r_metodos)) == len(metodos.METODOS)
    assert set(vm.asignacion_trimestral(r_metodos)["la_aportacion_va_a"]) <= {*r_metodos.paneles, "—"}


def test_el_pdf_se_arma(r_metodos):
    from src.export.pdf_metodos import html_de_metodos

    documento = html_de_metodos(r_metodos, fuentes_css="")
    for titulo in ("¿Predicen?", "A cuál de los tres va la aportación", "Metodología", "Qué hacer hoy"):
        assert titulo in documento


_SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


@pytest.mark.skipif(_SOFFICE is None, reason="LibreOffice no está instalado.")
def test_las_formulas_del_excel_reproducen_el_estudio(r_metodos, tmp_path):
    from openpyxl import load_workbook

    from src.export import excel_metodos
    from src.export.excel import buscar_errores

    ruta = excel_metodos.exportar(r_metodos, tmp_path / "valuacion_trimestral.xlsx")
    salida = tmp_path / "recalc"
    subprocess.run([_SOFFICE, "--headless", f"-env:UserInstallation=file://{tmp_path}/perfil",
                    "--convert-to", "xlsx", "--outdir", str(salida), str(ruta)],
                   capture_output=True, timeout=300, check=False)
    recalculado = salida / ruta.name
    assert recalculado.exists()
    assert buscar_errores(recalculado) == []
    ws = load_workbook(recalculado, data_only=True)["O"]
    encabezado = [c.value for c in ws[1]]
    ultimo = {h: c.value for h, c in zip(encabezado, ws[ws.max_row], strict=True)}
    q = metodos.trimestral(r_metodos.paneles["O"]).dropna(subset=["precio"]).iloc[-1]
    assert ultimo["Múltiplo precio/flujo"] == pytest.approx(q["multiplo_flujo"], rel=1e-9)
    assert ultimo["Prima sobre la tasa real"] == pytest.approx(q["v_prima_real"], abs=1e-12)
    assert ultimo["Prima sobre bonos Baa"] == pytest.approx(q["v_prima_baa"], abs=1e-12)
    assert ultimo["Señal del consenso"] == q["s_consenso"]
