"""Prueba 38 — El libro de Excel cableado a los estados.

Qué cambió
----------
Antes eran dos libros: uno con los estados financieros y otro con el modelo de
valuación. Se podían leer los dos, pero cambiar un renglón del balance no movía
nada del modelo, así que el estado era una lámina y no un insumo.

Ahora es uno solo, y los insumos que SALEN de un estado dejan de ser un número
escrito para volverse una fórmula que apunta a su renglón. Cambiar la deuda en el
balance mueve el apalancamiento, el LTV y el NAV.

Los dos errores que encontró hacerlo
------------------------------------
**El libro exportaba el balance en CEROS.** La pantalla armaba sus
``InsumosValuacion`` para exportar sin los cinco saldos ni los dos TTM, así que
el archivo que el usuario se llevaba traía deuda 0, efectivo 0, goodwill 0,
EBITDAre 0 e intereses 0 — y su hoja de Valuación calculaba apalancamiento, LTV y
NAV sobre nada. En la pantalla los números salían bien: el error solo existía en
el archivo (38.1).

**Cablear la deuda al renglón «Deuda total» revertía el PR #22.** Parecía lo
obvio y es lo equivocado: en Realty Income ese renglón trae 25,091.6 millones
—una sola etiqueta, que son sus notas senior— mientras que sus cuatro tramos
suman 30,651.7. Con el cableado ingenuo el libro devolvía la cifra vieja y el
apalancamiento bajaba de 5.68x a 4.63x, sin que nada lo dijera (38.4).

Ese segundo lo cazó **recalcular el libro con LibreOffice** y comparar celda por
celda contra el modelo. Un cableado que no se recalcula no está verificado: está
escrito. La comparación completa vive en 38.6 y se salta donde no hay
LibreOffice, así que 38.4 la sostiene también en Python puro.
"""

from __future__ import annotations

import datetime as dt
import importlib.util
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from src.datos.repositorio import Repositorio  # noqa: E402
from src.export.excel import (  # noqa: E402
    CABLEADOS_DE_BALANCE,
    CELDAS_INPUT,
    HOJA_PROPIA,
    VERDE,
    DatosExportacion,
    EstadosParaLibro,
    _formula_deuda,
    estados_para_libro,
    exportar,
)
from src.modelo.valuacion import InsumosValuacion  # noqa: E402
from src.servicio import construir_panel, saldos_de_balance  # noqa: E402

# Los siete insumos que la hoja de Inputs pide por renglón de balance o de TTM.
INSUMOS_DEL_BALANCE = (
    "deuda_total", "efectivo", "prestamos_por_cobrar",
    "inversiones_no_consolidadas", "goodwill", "ebitdare_ttm", "intereses_ttm",
)


def _insumos_reales(repo, ticker: str, asof: dt.date) -> InsumosValuacion:
    """Los mismos que arma la pantalla al exportar, después del arreglo."""
    panel = construir_panel(repo, ticker, asof=asof)
    sub = panel.trimestral.dropna(subset=["affo_por_accion_ttm"])
    fila = (sub if not sub.empty else panel.trimestral).tail(1).iloc[0]
    saldos = saldos_de_balance(repo, ticker, asof=asof)

    def numero(valor):
        try:
            return float(valor)
        except (TypeError, ValueError):
            return None

    return InsumosValuacion(
        ticker=ticker,
        precio=panel.precio or 0.0,
        acciones_diluidas=numero(fila.get("acciones_diluidas")) or 1.0,
        affo_ttm=numero(fila.get("affo_ttm")) or 0.0,
        affo_por_accion_ttm=numero(fila.get("affo_por_accion_ttm")),
        dividendo_ttm_por_accion=panel.dividendo_ttm,
        deuda_total=saldos.get("deuda_total") or 0.0,
        efectivo=saldos.get("efectivo") or 0.0,
        prestamos_por_cobrar=saldos.get("prestamos_por_cobrar") or 0.0,
        inversiones_no_consolidadas=saldos.get("inversiones_no_consolidadas") or 0.0,
        goodwill=saldos.get("goodwill") or 0.0,
        ebitdare_ttm=numero(fila.get("ebitdare_ttm")),
        intereses_ttm=numero(fila.get("gasto_intereses_ttm")),
        sector=panel.sector,
    )


@pytest.fixture(scope="module")
def libro(tmp_path_factory):
    """Un libro real de Realty Income, cableado. Se arma una vez."""
    repo = Repositorio()
    hoy = dt.date.today()
    panel = construir_panel(repo, "O", asof=hoy)
    if panel.trimestral.empty:
        pytest.skip("No hay base cargada.")
    datos = DatosExportacion(
        ticker="O", nombre="Realty Income", sector=panel.sector, fecha_corte=hoy,
        insumos=_insumos_reales(repo, "O", hoy), componentes_cascada={},
        estados=estados_para_libro(repo, "O", asof=hoy, n_periodos=8),
    )
    ruta = tmp_path_factory.mktemp("libro") / "O.xlsx"
    exportar(datos, ruta)
    return ruta


# --------------------------------------------------------------------------------------
# 38.1 · Un libro, no dos — y con el balance lleno
# --------------------------------------------------------------------------------------


def test_el_libro_trae_los_estados_y_el_modelo(libro):
    hojas = load_workbook(libro).sheetnames
    assert "Inputs" in hojas and "Valuación" in hojas, hojas
    assert HOJA_PROPIA["balance"] in hojas, hojas
    assert any(h.startswith("BBG") for h in hojas), hojas
    assert any(h.startswith("Reportado") for h in hojas), hojas
    assert len(set(hojas)) == len(hojas) and all(len(h) <= 31 for h in hojas)


def test_los_insumos_del_balance_ya_no_salen_en_cero():
    """El error que solo existía en el archivo.

    La pantalla mostraba el apalancamiento correcto y el libro exportado lo
    calculaba sobre una deuda de cero, un efectivo de cero y un EBITDAre de cero.
    Nada en la pantalla lo delataba porque el error no estaba en la pantalla.
    """
    repo = Repositorio()
    if construir_panel(repo, "O", asof=dt.date.today()).trimestral.empty:
        pytest.skip("No hay base cargada.")
    ins = _insumos_reales(repo, "O", dt.date.today())
    assert ins.deuda_total > 1e9, "la deuda volvió a salir en cero"
    assert ins.efectivo > 0
    assert ins.ebitdare_ttm and ins.ebitdare_ttm > 0
    assert ins.intereses_ttm and ins.intereses_ttm > 0


# --------------------------------------------------------------------------------------
# 38.2 · Lo que sale de un estado es fórmula; lo que es supuesto, no
# --------------------------------------------------------------------------------------


def test_los_insumos_de_balance_son_formulas_verdes(libro):
    ws = load_workbook(libro)["Inputs"]
    for clave in ("deuda_total", "efectivo", "goodwill", "intereses_ttm"):
        celda = ws[CELDAS_INPUT[clave]]
        assert str(celda.value).startswith("="), f"{clave} sigue siendo un número pegado"
        assert celda.font.color.rgb == VERDE, f"{clave} no está marcado como referencia"


def test_los_supuestos_siguen_siendo_del_usuario(libro):
    """El precio, el cap rate y las tasas NO salen de un estado financiero.

    Volverlos fórmula habría escondido la única parte del modelo que de verdad es
    del usuario, y es justo la que la hoja de Sensibilidad existe para mover.
    """
    ws = load_workbook(libro)["Inputs"]
    for clave in ("precio", "cap_rate_mercado", "tasa_libre_riesgo",
                  "yield_adquisiciones", "peso_deuda_marginal"):
        celda = ws[CELDAS_INPUT[clave]]
        assert not str(celda.value).startswith("="), f"{clave} dejó de ser editable"
        assert celda.fill.fgColor.rgb == "FFFFFF99", f"{clave} perdió el amarillo"


# --------------------------------------------------------------------------------------
# 38.3 · La referencia va por etiqueta, no por número de fila
# --------------------------------------------------------------------------------------


def test_las_referencias_buscan_el_renglon_por_su_nombre(libro):
    """Lo que hace que el cableado sobreviva a la próxima exportación.

    Un estado no tiene forma fija: si la emisora empieza a reportar un renglón
    que antes no tenía, todo lo que va debajo se recorre una fila. Una referencia
    dura seguiría apuntando a la fila, que ahora es otra partida, y el modelo
    cambiaría de insumo sin decirlo. Con MATCH, un renglón que desaparece da
    #N/A: un error visible, que es infinitamente mejor.
    """
    ws = load_workbook(libro)["Inputs"]
    for clave in ("efectivo", "goodwill", "deuda_total"):
        formula = str(ws[CELDAS_INPUT[clave]].value)
        assert "MATCH(" in formula and "INDEX(" in formula, clave


def test_ninguna_formula_apunta_a_una_celda_dura_del_estado(libro):
    """Una referencia como 'Propia Balance'!I14 es exactamente lo prohibido."""
    import re

    ws = load_workbook(libro)["Inputs"]
    duras = []
    for clave in INSUMOS_DEL_BALANCE:
        formula = str(ws[CELDAS_INPUT[clave]].value)
        if re.search(r"'Propia [^']+'![A-Z]+\d+", formula):
            duras.append(clave)
    assert duras == [], f"apuntan a una fila fija: {duras}"


# --------------------------------------------------------------------------------------
# 38.4 · La deuda son los TRAMOS, no el renglón declarado
# --------------------------------------------------------------------------------------


def test_la_deuda_no_se_cablea_al_renglon_declarado(libro):
    """El error que revertía el PR #22, cazado por la recalculación.

    El renglón «Deuda total» de Realty Income trae 25,091.6 millones porque sale
    de una sola etiqueta —sus notas senior—. Sus tramos suman 30,651.7. Cablear
    el renglón bajaba el apalancamiento de 5.68x a 4.63x en silencio.
    """
    from src.ingesta.estados import LINEA_POR_CLAVE
    from src.servicio import TRAMOS_DE_DEUDA

    formula = str(load_workbook(libro)["Inputs"][CELDAS_INPUT["deuda_total"]].value)
    etiquetas = {LINEA_POR_CLAVE[c].etiqueta for c in TRAMOS_DE_DEUDA
                 if c in LINEA_POR_CLAVE}
    citados = sum(1 for e in etiquetas if f'"{e}"' in formula)
    assert citados >= 2, f"solo cita {citados} tramo(s): {formula[:200]}"
    assert "deuda_total" not in {c for c, _ in CABLEADOS_DE_BALANCE}, (
        "la deuda volvió a la lista de cableados simples"
    )


def test_la_suma_de_tramos_solo_gana_si_supera_el_margen():
    """Un redondeo no cambia la fuente: la regla es la del servicio."""
    from src.servicio import MARGEN_DE_TRAMOS

    tabla = pd.DataFrame({
        "Renglón": ["Deuda total", "Notas senior no garantizadas", "Deuda hipotecaria",
                    "Pasivos totales"],
        "2026-06-30": [100.0, 60.0, 40.0, 500.0],
    })
    formula = _formula_deuda("Propia Balance", tabla, "B", 1.0)
    assert f"*{1 + MARGEN_DE_TRAMOS}" in formula, formula
    assert "MIN(" in formula, "no topa con los pasivos totales"


def test_sin_tramos_declarados_no_se_inventa_una_deuda():
    tabla = pd.DataFrame({"Renglón": ["Efectivo y equivalentes"], "2026-06-30": [10.0]})
    assert _formula_deuda("Propia Balance", tabla, "B", 1.0) == ""


# --------------------------------------------------------------------------------------
# 38.5 · La guarda de la depreciación, dentro de Excel
# --------------------------------------------------------------------------------------


def test_el_ebitdare_del_libro_lleva_la_guarda_de_la_depreciacion(libro):
    """La prueba 36, una capa más afuera.

    En Python, sumar con la depreciación ausente daba un EBITDAre que no lo era.
    En Excel el riesgo es peor: ``SUM`` de una celda vacía **es cero** y no avisa.
    Por eso la fórmula cuenta primero cuántos trimestres traen depreciación.
    """
    formula = str(load_workbook(libro)["Inputs"][CELDAS_INPUT["ebitdare_ttm"]].value)
    assert formula.startswith("=IF("), formula[:80]
    assert "ISNUMBER(" in formula, "no verifica que la depreciación exista"
    assert '"")' in formula or ',"",' in formula, "no devuelve vacío cuando falta"
    assert "Depreciación y amortización" in formula


# --------------------------------------------------------------------------------------
# 38.6 · Recalculado de verdad: contra el modelo, celda por celda
# --------------------------------------------------------------------------------------

_SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
# Leer el libro RECALCULADO exige odfpy: LibreOffice guarda el resultado de cada
# fórmula en el .ods, y openpyxl solo sabe leer el .xlsx, donde openpyxl mismo
# escribió la fórmula sin resultado. Sin las dos piezas la prueba no puede correr,
# y decirlo es mejor que reventar con un ImportError a media suite.
_ODF = importlib.util.find_spec("odf") is not None


@pytest.mark.skipif(
    _SOFFICE is None or not _ODF,
    reason="Se necesitan LibreOffice y odfpy para recalcular y leer el resultado.",
)
def test_el_libro_recalculado_da_lo_mismo_que_el_modelo(libro, tmp_path):
    """La verificación que cazó el error de la deuda.

    Las demás pruebas leen la FÓRMULA; esta lee su RESULTADO. Es la diferencia
    entre "la fórmula menciona los tramos" y "la fórmula da 30,651.65".
    """
    repo = Repositorio()
    ins = _insumos_reales(repo, "O", dt.date.today())
    subprocess.run(
        [_SOFFICE, "--headless", "--calc", "--convert-to", "ods",
         "--outdir", str(tmp_path), str(libro)],
        check=True, capture_output=True, timeout=600,
    )
    recalculado = tmp_path / "O.ods"
    assert recalculado.exists(), "LibreOffice no produjo el libro recalculado"

    inputs = pd.read_excel(recalculado, sheet_name="Inputs", header=None, engine="odf")
    diferencias = []
    for clave in INSUMOS_DEL_BALANCE:
        fila = int(CELDAS_INPUT[clave][1:])
        crudo = inputs.iat[fila - 1, 1]
        esperado = getattr(ins, clave, None)
        if esperado is None or pd.isna(esperado):
            continue  # el modelo tampoco lo tiene: la guarda hizo lo suyo
        obtenido = float(crudo) if pd.notna(crudo) and not isinstance(crudo, str) else None
        objetivo = esperado / 1e6
        if obtenido is None or abs(obtenido - objetivo) > max(0.02, abs(objetivo) * 0.0005):
            diferencias.append(f"{clave}: libro {obtenido} vs modelo {objetivo:,.2f}")
    assert diferencias == [], "el libro no reproduce el modelo:\n" + "\n".join(diferencias)


# --------------------------------------------------------------------------------------
# 38.7 · Sin estados, el libro sigue siendo el de antes
# --------------------------------------------------------------------------------------


def test_un_libro_sin_estados_no_se_rompe(tmp_path):
    """La integración es opcional: quitarla no puede tumbar la exportación."""
    datos = DatosExportacion(
        ticker="X", nombre="Prueba", sector="net_lease", fecha_corte=dt.date(2026, 6, 30),
        insumos=InsumosValuacion(ticker="X", precio=50.0, acciones_diluidas=1e6),
        componentes_cascada={},
    )
    ruta = exportar(datos, tmp_path / "X.xlsx")
    hojas = load_workbook(ruta).sheetnames
    assert "Inputs" in hojas and "Valuación" in hojas
    assert not any(h.startswith(("BBG", "Propia", "Reportado")) for h in hojas)


def test_unas_vistas_vacias_no_dejan_referencias_rotas(tmp_path):
    """Si el balance viene vacío, el input se queda como estaba y no apunta a nada."""
    datos = DatosExportacion(
        ticker="X", nombre="Prueba", sector="net_lease", fecha_corte=dt.date(2026, 6, 30),
        insumos=InsumosValuacion(ticker="X", precio=50.0, acciones_diluidas=1e6,
                                 deuda_total=1_000_000.0),
        componentes_cascada={},
        estados=EstadosParaLibro(propia={}, bloomberg={}, reportados={},
                                 ratios=pd.DataFrame()),
    )
    ruta = exportar(datos, tmp_path / "X.xlsx")
    ws = load_workbook(ruta)["Inputs"]
    assert not str(ws[CELDAS_INPUT["deuda_total"]].value).startswith("=")
