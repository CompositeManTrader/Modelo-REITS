"""Prueba 39 — Fórmulas vivas y la historia completa.

Qué cambió
----------
Las dos cosas que el libro exportado todavía no hacía.

**Los cálculos venían pegados como número.** El margen de EBITDAre, la cobertura
de intereses, el apalancamiento y los ratios de la plantilla de Bloomberg
llegaban al Excel ya resueltos. Eran un retrato del modelo, no el modelo: mover
la depreciación en «Propia Resultados» no movía nada. Ahora cada renglón que se
CALCULA es una fórmula que apunta a los renglones de los que sale, y los únicos
números pegados son los que de verdad son dato —el hecho reportado y el no-GAAP
del comunicado—.

**El libro traía unos cuantos trimestres.** La exportación se recortaba a ocho.
Un modelo de valuación al que se le puede ver el ciclo entero necesita el ciclo
entero: ``n_periodos=None`` trae toda la historia —setenta trimestres en Realty
Income, dieciocho ejercicios— y la pantalla deja elegir.

Lo que encontró escribirlo
--------------------------
Traducir a Excel no es transcribir: es reimplementar, y una reimplementación se
separa del original en silencio. Las cinco formas en que se separó, todas
cazadas recalculando con LibreOffice y comparando celda por celda:

1. ``N()`` de una celda vacía **es cero**. La fórmula que suma sin verificar
   presencia daba número donde el modelo daba hueco: 57 celdas en la primera
   vuelta.
2. ``SUM`` de una ventana TTM incompleta suma lo que haya. Dos trimestres de
   EBITDA contra cuatro de ingresos daban márgenes de 376%.
3. Un crecimiento anual medido contra una base no positiva da −183% y parece un
   desplome.
4. ``(a)/(b)*4`` **no es** ``a/(b*4)``. El apalancamiento salía dieciséis veces
   el que es, y un 74.7x en una hoja de Excel no se distingue a simple vista de
   un 4.67x mal puesto.
5. La primera hoja de Ratios reescribió los cálculos en vez de traducirlos: el
   NOI salía del ingreso TOTAL y no del de renta —el error de Welltower, el que
   da un cap rate de 0.89%—, el EBITDAre exigía impuestos, y el apalancamiento
   no anualizaba. 356 de 923 celdas de Realty Income no cuadraban. Ninguna daba
   error.

De ahí la forma de estas pruebas: **una declaración, dos traductores**. El
cálculo se declara una vez —`FormulaRatio`, `InsumoDerivado`, `Formula` del molde
de Bloomberg— y de ahí salen el valor de la pantalla y la fórmula del libro. Y
39.5 lo comprueba donde importa: recalculando el libro y comparando contra
Python, celda por celda.
"""

from __future__ import annotations

import datetime as dt
import importlib.util
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from src.datos.repositorio import Repositorio  # noqa: E402
from src.export.excel import (  # noqa: E402
    TRIMESTRES_TTM,
    DatosExportacion,
    estados_para_libro,
    exportar,
)
from src.modelo import bloomberg as B  # noqa: E402
from src.modelo.valuacion import InsumosValuacion  # noqa: E402
from src.servicio import (  # noqa: E402
    FORMULA_RATIO,
    INSUMO_POR_CLAVE,
    INSUMOS_DE_RATIOS,
    RATIOS_PROPIOS,
    _derivar_ebitdare,
    _derivar_noi,
    panel_de_conceptos,
    ratios_propios,
    serie_de_insumo,
)

HOY = dt.date.today()


@pytest.fixture(scope="module")
def panel():
    p = panel_de_conceptos(Repositorio(), "O", asof=HOY)
    if p.empty:
        pytest.skip("No hay base cargada.")
    return p


@pytest.fixture(scope="module")
def libro(tmp_path_factory):
    """Un libro real de Realty Income, con TODA su historia."""
    repo = Repositorio()
    estados = estados_para_libro(repo, "O", asof=HOY)
    if not estados.hay_propia:
        pytest.skip("No hay base cargada.")
    datos = DatosExportacion(
        ticker="O", nombre="Realty Income", sector="net_lease", fecha_corte=HOY,
        insumos=InsumosValuacion(ticker="O", precio=60.0, acciones_diluidas=9e8),
        componentes_cascada={}, estados=estados,
    )
    return exportar(datos, tmp_path_factory.mktemp("vivas") / "O.xlsx")


# --------------------------------------------------------------------------------------
# 39.1 · Una declaración, no dos implementaciones
# --------------------------------------------------------------------------------------


def test_cada_ratio_de_la_pantalla_tiene_su_declaracion():
    """Un ratio sin `FormulaRatio` sale en la pantalla y vacío en el libro."""
    faltan = [c for _, c, _, _ in RATIOS_PROPIOS if c not in FORMULA_RATIO]
    assert faltan == [], faltan


def test_las_claves_con_arroba_son_insumos_declarados():
    """La arroba dice «esto no es un renglón»; tiene que apuntar a un insumo."""
    for clave, formula in FORMULA_RATIO.items():
        for _signo, termino in formula.numerador + formula.denominador:
            if termino.startswith("@"):
                assert termino in INSUMO_POR_CLAVE, f"{clave} usa {termino}"


def test_el_insumo_del_noi_reproduce_la_regla_del_modelo(panel):
    """El NOI del libro es el del modelo, no una reescritura parecida.

    La reescritura existió: derivaba del ingreso TOTAL en vez del de renta. Es
    exactamente el error que `_derivar_noi` documenta —Welltower factura por
    operación de vivienda, no por renta triple neta— y no levanta excepción.
    """
    declarado = serie_de_insumo(panel, INSUMO_POR_CLAVE["@noi"])
    del_modelo = _derivar_noi(panel.drop(columns=["noi"], errors="ignore"))
    pd.testing.assert_series_equal(
        declarado, del_modelo, check_names=False, rtol=1e-12
    )


def test_el_insumo_del_ebitdare_reproduce_la_regla_del_modelo(panel):
    """Y el EBITDAre también, con la guarda de la depreciación incluida."""
    declarado = serie_de_insumo(panel, INSUMO_POR_CLAVE["@ebitdare"])
    del_modelo = _derivar_ebitdare(panel)
    pd.testing.assert_series_equal(
        declarado, del_modelo, check_names=False, rtol=1e-12
    )


def test_sin_depreciacion_no_hay_ebitdare_declarado():
    """La prueba 36, sobre la declaración: un cero ahí es otro número, no un hueco."""
    panel = pd.DataFrame(
        [{"utilidad_neta": 463e6, "gasto_intereses": 182e6, "impuestos": -62e6,
          "depreciacion_amortizacion": None, "ingresos_totales": 3_545e6}],
        index=pd.DatetimeIndex([dt.date(2026, 6, 30)]),
    )
    assert pd.isna(serie_de_insumo(panel, INSUMO_POR_CLAVE["@ebitdare"]).iloc[0])


def test_los_opcionales_valen_cero_y_los_exigidos_no():
    """La diferencia entre un dato ausente y un cero legítimo, medida.

    Un trimestre sin deterioro es el trimestre normal. Un trimestre sin
    depreciación es un trimestre incompleto. Tratarlos igual —en cualquiera de
    las dos direcciones— cuesta media docena de emisoras o un número inventado.
    """
    base = {"utilidad_neta": 100.0, "gasto_intereses": 20.0,
            "depreciacion_amortizacion": 50.0, "deterioro": None, "impuestos": None}
    panel = pd.DataFrame([base], index=pd.DatetimeIndex([dt.date(2026, 6, 30)]))
    assert serie_de_insumo(panel, INSUMO_POR_CLAVE["@ebitdare"]).iloc[0] == 170.0

    sin_intereses = dict(base, gasto_intereses=0.0)
    panel = pd.DataFrame([sin_intereses], index=pd.DatetimeIndex([dt.date(2026, 6, 30)]))
    assert pd.isna(serie_de_insumo(panel, INSUMO_POR_CLAVE["@ebitdare"]).iloc[0]), (
        "sin intereses no es EBITDAre: es utilidad operativa con otro nombre"
    )


def test_el_apalancamiento_anualiza_el_denominador_y_el_costo_el_numerador():
    """Los dos lados anualizan, y no es simetría decorativa.

    El costo de la deuda compara intereses del TRIMESTRE contra la deuda al
    corte; el apalancamiento compara la deuda al corte contra el EBITDAre del
    TRIMESTRE. Con un solo campo para las dos, uno de los dos sale con un factor
    de cuatro de error.
    """
    assert FORMULA_RATIO["costo_deuda"].anualiza_numerador == 4
    assert FORMULA_RATIO["costo_deuda"].anualiza_denominador == 1
    assert FORMULA_RATIO["apalancamiento"].anualiza_numerador == 1
    assert FORMULA_RATIO["apalancamiento"].anualiza_denominador == 4

    panel = pd.DataFrame(
        [{"deuda_total": 1_000.0, "efectivo": 100.0, "utilidad_neta": 30.0,
          "gasto_intereses": 10.0, "depreciacion_amortizacion": 60.0}],
        index=pd.DatetimeIndex([dt.date(2026, 6, 30)]),
    )
    tabla = ratios_propios(panel)
    fila = tabla[tabla["Ratio"] == "Deuda neta / EBITDAre"].iloc[0]
    # (1000 − 100) ÷ (100 × 4) = 2.25
    assert fila["2026-06-30"] == pytest.approx(2.25)


def test_el_efectivo_es_opcional_y_la_deuda_no():
    """La deuda neta sin efectivo es la deuda; sin deuda no hay deuda neta."""
    sin_efectivo = pd.DataFrame(
        [{"deuda_total": 1_000.0, "utilidad_neta": 30.0, "gasto_intereses": 10.0,
          "depreciacion_amortizacion": 60.0}],
        index=pd.DatetimeIndex([dt.date(2026, 6, 30)]),
    )
    tabla = ratios_propios(sin_efectivo)
    fila = tabla[tabla["Ratio"] == "Deuda neta / EBITDAre"].iloc[0]
    assert fila["2026-06-30"] == pytest.approx(2.5)

    sin_deuda = sin_efectivo.drop(columns=["deuda_total"])
    tabla = ratios_propios(sin_deuda)
    fila = tabla[tabla["Ratio"] == "Deuda neta / EBITDAre"].iloc[0]
    assert pd.isna(fila["2026-06-30"])


# --------------------------------------------------------------------------------------
# 39.2 · Las guardas que Excel necesita y Python no
# --------------------------------------------------------------------------------------


def _formulas(ws) -> list[str]:
    return [str(c.value) for fila in ws.iter_rows() for c in fila
            if isinstance(c.value, str) and c.value.startswith("=")]


def test_ninguna_suma_del_libro_confia_en_una_celda_vacia(libro):
    """``N()`` de una celda vacía es cero, y un cero ahí no es un hueco.

    Toda fórmula que suma renglones tiene que llevar su guarda. Hay dos, y cuál
    va es parte de la declaración: PRESENCIA —``COUNT``, ``ISNUMBER``— cuando el
    dato se exige en ese periodo, y POSITIVIDAD —``<=0``— cuando el hueco cuenta
    como cero y lo que descalifica es el resultado. Sin ninguna de las dos, el
    libro daba número donde el modelo daba vacío: 57 celdas en la primera vuelta,
    y ninguna marcaba nada.
    """
    wb = load_workbook(libro)
    for nombre in ("BBG Resultados", "BBG Balance", "BBG Flujo", "Ratios"):
        for formula in _formulas(wb[nombre]):
            if "N(" not in formula:
                continue
            assert ("COUNT(" in formula or "ISNUMBER(" in formula
                    or "<=0" in formula), (
                f"{nombre}: suma sin guarda — {formula[:120]}"
            )


def test_cada_insumo_lleva_la_guarda_que_declaro(libro):
    """Y la que declaró, no la otra: son reglas distintas, no dos estilos.

    El EBITDAre exige el dato EN ESE PERIODO —la depreciación en cero da otro
    número, prueba 36—. El NOI no: un trimestre sin gasto del inmueble cuenta
    como cero y lo que descalifica es que el resultado no sea positivo o no
    llegue al piso sobre el ingreso. Intercambiarlas borra emisoras enteras en un
    sentido e inventa números en el otro.
    """
    ws = load_workbook(libro)["Ratios"]
    for insumo in INSUMOS_DE_RATIOS:
        fila = next((f[0].row for f in ws.iter_rows(min_col=1, max_col=1)
                     if str(f[0].value) == insumo.etiqueta), None)
        assert fila, insumo.etiqueta
        formula = next(str(c.value) for c in ws[fila]
                       if isinstance(c.value, str) and c.value.startswith("="))
        if insumo.por_celda:
            assert "ISNUMBER(" in formula, f"{insumo.etiqueta}: {formula[:120]}"
        else:
            assert "<=0" in formula, f"{insumo.etiqueta}: {formula[:120]}"
            assert insumo.piso is None or f"<{insumo.piso[0]}" in formula, (
                f"{insumo.etiqueta}: sin el piso — {formula[:160]}"
            )


def test_ninguna_ventana_ttm_se_conforma_con_lo_que_haya(libro):
    """Doce meses son doce meses: ``SUM`` de una ventana incompleta miente."""
    wb = load_workbook(libro)
    for nombre in ("BBG Resultados", "BBG Balance", "BBG Flujo"):
        for formula in _formulas(wb[nombre]):
            if "SUM(" not in formula:
                continue
            assert "COUNT(" in formula and f"<{TRIMESTRES_TTM}" in formula, (
                f"{nombre}: TTM sin exigir los cuatro trimestres — {formula[:120]}"
            )


def test_ningun_crecimiento_se_mide_contra_una_base_no_positiva(libro):
    """Un crecimiento contra un FFO negativo da −183% y parece un desplome."""
    wb = load_workbook(libro)
    yoy = [ln.etiqueta for ln in B.PLANTILLA[B.RESULTADOS]
           if ln.formula is not None and ln.formula.yoy]
    if not yoy:
        pytest.skip("El molde no declara crecimientos anuales.")
    ws = wb["BBG Resultados"]
    encontradas = 0
    for fila in ws.iter_rows(min_col=1, max_col=1):
        if str(fila[0].value) not in yoy:
            continue
        for celda in ws[fila[0].row]:
            if isinstance(celda.value, str) and celda.value.startswith("="):
                assert "<=0" in celda.value, celda.value[:120]
                encontradas += 1
    assert encontradas, "no se escribió ningún crecimiento anual"


def _denominador(formula: str) -> str:
    """El divisor de una fórmula ``=IF(OR(...),"",num/den)``, como texto."""
    cuerpo = formula[formula.index(',"",') + 4: -1]
    nivel = 0
    for i, caracter in enumerate(cuerpo):
        nivel += (caracter == "(") - (caracter == ")")
        if caracter == "/" and nivel == 0:
            return cuerpo[i + 1:]
    raise AssertionError("la fórmula no divide: " + formula[:160])


def test_el_apalancamiento_del_libro_divide_entre_el_denominador_anualizado(libro):
    """``(a)/(b)*4`` no es ``a/(b*4)``, y la diferencia es de dieciséis veces.

    Escrito sin el paréntesis, el apalancamiento de Realty Income salía en 74.7x
    donde son 4.67x. Es un número que un lector cuidadoso rechaza a la primera
    —ningún REIT vive a 74 veces— pero que ninguna celda marca como error, y en
    57 columnas seguidas se vuelve paisaje.
    """
    ws = load_workbook(libro)["Ratios"]
    fila = next(f[0].row for f in ws.iter_rows(min_col=1, max_col=1)
                if str(f[0].value) == "Deuda neta / EBITDAre")
    formula = next(str(c.value) for c in ws[fila]
                   if isinstance(c.value, str) and c.value.startswith("="))
    den = _denominador(formula)
    assert "*4" in den, "no anualizó el denominador: " + formula[:160]
    assert den.startswith("(") and den.endswith(")") and not den.endswith("*4"), (
        "el factor quedó FUERA del denominador: " + den[-40:]
    )


# --------------------------------------------------------------------------------------
# 39.3 · Los ratios, como fórmula y no como retrato
# --------------------------------------------------------------------------------------


def test_la_hoja_de_ratios_no_trae_un_solo_numero_calculado(libro):
    """Lo pegado tiene que ser dato; lo calculado, fórmula.

    El no-GAAP del comunicado —AFFO, FFO— va como valor porque es dato: lo
    publica la emisora, no sale de sumar renglones del 10-Q. Todo lo demás es
    fórmula, y por eso escribir otro AFFO mueve el payout.
    """
    ws = load_workbook(libro)["Ratios"]
    etiquetas_de_dato = {"FFO", "AFFO", "AFFO por acción",
                         "Del comunicado (no-GAAP)", "Insumos derivados"}
    calculadas = {ln.etiqueta for ln in INSUMOS_DE_RATIOS} | {
        e for e, _, _, _ in RATIOS_PROPIOS
    }
    # Solo las columnas de PERIODO: la última trae la explicación en prosa.
    ultima = max(c.column for c in ws[1] if c.value and c.value != "Se calcula así")
    vistas = 0
    for fila in ws.iter_rows(min_row=2, max_col=ultima):
        etiqueta = str(fila[0].value)
        if etiqueta not in calculadas:
            assert etiqueta in etiquetas_de_dato or etiqueta in ("Ratios", "None"), etiqueta
            continue
        vistas += 1
        for celda in fila[1:]:
            if celda.value is None or isinstance(celda.value, str) and not celda.value:
                continue
            assert isinstance(celda.value, str) and celda.value.startswith("="), (
                f"{etiqueta}: número pegado — {celda.value!r}"
            )
    assert vistas == len(calculadas), f"faltaron renglones: {vistas} de {len(calculadas)}"


def test_los_ratios_apuntan_a_los_renglones_del_estado(libro):
    """Y apuntan POR ETIQUETA: un renglón nuevo no puede cambiarles el insumo."""
    ws = load_workbook(libro)["Ratios"]
    formulas = [f for f in _formulas(ws) if "Propia" in f]
    assert formulas, "ningún ratio referencia un estado"
    for formula in formulas:
        assert "MATCH(" in formula, formula[:120]


def test_los_ratios_no_abren_columna_para_un_corte_sin_resultados(libro):
    """El balance trae cortes que ningún estado de resultados acompaña.

    Una columna de ratios sobre ese corte es la deuda sola, y se lee como si al
    trimestre le faltara todo lo demás.
    """
    wb = load_workbook(libro)
    def periodos(hoja):
        ws = wb[hoja]
        return {str(c.value) for c in ws[1][1:] if c.value}
    de_ratios = periodos("Ratios") - {"Se calcula así"}
    de_resultados = periodos("Propia Resultados") - {"Renglón"}
    de_flujo = periodos("Propia Flujo") - {"Renglón"} if "Propia Flujo" in wb else set()
    assert de_ratios <= (de_resultados | de_flujo), sorted(
        de_ratios - de_resultados - de_flujo
    )[:5]


# --------------------------------------------------------------------------------------
# 39.4 · Toda la historia, no unos cuantos trimestres
# --------------------------------------------------------------------------------------


def test_el_panel_sin_recorte_trae_mas_que_los_ocho_de_antes(panel):
    assert len(panel) > 8, f"solo {len(panel)} periodos"


def test_el_libro_exporta_los_periodos_que_el_panel_tiene(libro, panel):
    """Lo que se descarga es la historia, no una ventana."""
    ws = load_workbook(libro)["Propia Resultados"]
    columnas = [c.value for c in ws[1][1:] if c.value]
    assert len(columnas) >= len(panel) - 2, (
        f"el libro trae {len(columnas)} periodos y el panel {len(panel)}"
    )


def test_el_recorte_sigue_disponible_para_quien_lo_pida():
    """La pantalla deja elegir; el recorte no desapareció, dejó de ser el default."""
    repo = Repositorio()
    corto = panel_de_conceptos(repo, "O", asof=HOY, n_periodos=8)
    if corto.empty:
        pytest.skip("No hay base cargada.")
    assert len(corto) == 8


# --------------------------------------------------------------------------------------
# 39.5 · Recalculado: la fórmula da lo mismo que Python
# --------------------------------------------------------------------------------------

_SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
_ODF = importlib.util.find_spec("odf") is not None


@pytest.mark.skipif(
    _SOFFICE is None or not _ODF,
    reason="Se necesitan LibreOffice y odfpy para recalcular y leer el resultado.",
)
def test_la_hoja_de_ratios_recalculada_da_lo_mismo_que_python(libro, panel, tmp_path):
    """La verificación que caza las reescrituras. Las demás leen la FÓRMULA.

    Esta lee su RESULTADO, que es la diferencia entre «la fórmula menciona el
    EBITDAre» y «la fórmula da 4.67x». Es la que encontró las cinco separaciones
    del encabezado de este archivo.
    """
    subprocess.run(
        [_SOFFICE, "--headless", "--calc", "--convert-to", "ods",
         "--outdir", str(tmp_path), str(libro)],
        check=True, capture_output=True, timeout=900,
    )
    recalculado = tmp_path / "O.ods"
    assert recalculado.exists(), "LibreOffice no produjo el libro recalculado"

    hoja = pd.read_excel(recalculado, sheet_name="Ratios", header=0, engine="odf")
    periodos = [str(c) for c in hoja.columns[1:] if not str(c).startswith("Unnamed")]
    fila_de = {}
    for i in range(len(hoja)):
        etiqueta = hoja.iloc[i, 0]
        if isinstance(etiqueta, str) and etiqueta:
            fila_de.setdefault(etiqueta, i)

    esperado: dict[str, dict[str, float]] = {}
    for insumo in INSUMOS_DE_RATIOS:
        serie = serie_de_insumo(panel, insumo)
        esperado[insumo.etiqueta] = {
            i.date().isoformat(): (None if pd.isna(v) else float(v))
            for i, v in serie.items()
        }
    for _, fila in ratios_propios(panel).iterrows():
        esperado[str(fila["Ratio"])] = {
            k: v for k, v in fila.items()
            if k not in ("Ratio", "formato", "explicacion")
        }

    diferencias, comparadas = [], 0
    for etiqueta, valores in esperado.items():
        assert etiqueta in fila_de, f"«{etiqueta}» no está en la hoja"
        i = fila_de[etiqueta]
        for j, periodo in enumerate(periodos):
            v_py = valores.get(periodo)
            v_xl = hoja.iloc[i, 1 + j]
            hay_py = v_py is not None and pd.notna(v_py)
            hay_xl = pd.notna(v_xl) and not isinstance(v_xl, str)
            comparadas += 1
            if not hay_py and not hay_xl:
                continue
            if hay_py != hay_xl:
                diferencias.append(f"{etiqueta}@{periodo}: py={v_py} libro={v_xl}")
            elif abs(float(v_xl) - float(v_py)) > max(1e-9, abs(float(v_py)) * 1e-6):
                diferencias.append(
                    f"{etiqueta}@{periodo}: py={float(v_py):.8g} libro={float(v_xl):.8g}"
                )
    assert comparadas > 100, f"solo se compararon {comparadas} celdas"
    assert diferencias == [], (
        f"{len(diferencias)} de {comparadas} celdas no cuadran:\n"
        + "\n".join(diferencias[:10])
    )


@pytest.mark.skipif(
    _SOFFICE is None or not _ODF,
    reason="Se necesitan LibreOffice y odfpy para recalcular y leer el resultado.",
)
def test_las_hojas_de_bloomberg_recalculadas_dan_lo_mismo_que_python(
    libro, panel, tmp_path
):
    """Lo mismo para el molde de Bloomberg: sus ratios también son fórmula."""
    subprocess.run(
        [_SOFFICE, "--headless", "--calc", "--convert-to", "ods",
         "--outdir", str(tmp_path), str(libro)],
        check=True, capture_output=True, timeout=900,
    )
    recalculado = tmp_path / "O.ods"
    hojas = {"BBG Resultados": B.RESULTADOS, "BBG Balance": B.BALANCE,
             "BBG Flujo": B.FLUJO}
    diferencias, comparadas = [], 0
    for nombre, estado in hojas.items():
        py = B.armar(panel, estado)
        periodos = [c for c in py.columns
                    if c != "Renglón" and c not in B.COLUMNAS_DE_APOYO]
        xl = pd.read_excel(recalculado, sheet_name=nombre, header=0, engine="odf")
        for i, linea in enumerate(B.PLANTILLA[estado]):
            if linea.formula is None:
                continue
            for j, periodo in enumerate(periodos):
                v_py = py.iloc[i][periodo]
                v_xl = xl.iloc[i, 1 + j]
                hay_py = v_py is not None and pd.notna(v_py)
                hay_xl = pd.notna(v_xl) and not isinstance(v_xl, str)
                comparadas += 1
                if not hay_py and not hay_xl:
                    continue
                if hay_py != hay_xl:
                    diferencias.append(
                        f"{nombre}/{linea.etiqueta}@{periodo}: py={v_py} libro={v_xl}"
                    )
                elif abs(float(v_xl) - float(v_py)) > max(1e-6, abs(float(v_py)) * 1e-6):
                    diferencias.append(
                        f"{nombre}/{linea.etiqueta}@{periodo}: "
                        f"py={float(v_py):.8g} libro={float(v_xl):.8g}"
                    )
    assert comparadas > 100, f"solo se compararon {comparadas} celdas"
    assert diferencias == [], (
        f"{len(diferencias)} de {comparadas} celdas no cuadran:\n"
        + "\n".join(diferencias[:10])
    )
