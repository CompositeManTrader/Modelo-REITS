"""El compilado: la descomposición del retorno, las preguntas y la calculadora del Excel."""

from __future__ import annotations

import numpy as np
import openpyxl
import pytest

from src.export import compilado as co
from src.export import excel_compilado as xc
from src.export.excel import RELLENO_INPUT
from src.investigacion import resultados


@pytest.mark.parametrize("ticker", ["O", "NNN"])
def test_la_descomposicion_cuadra(ticker):
    d = co.descomposicion(ticker)
    producto = (1 + d["dividendo"]) * (1 + d["crecimiento"]) * (1 + d["revaluacion"])
    assert producto == pytest.approx(1 + d["retorno"])
    assert 0 < d["dividendo"] < d["retorno"]                      # el dividendo es la mayor parte, no todo
    assert d["yield_hoy"] < d["yield_inicio"]                     # hoy pagan menos sobre su precio que al inicio


def test_anios_para_llegar_a_una_meta():
    anios = co._anios_para(500, 60_000, 0.07)
    i = 1.07 ** (1 / 12) - 1
    n = anios * 12
    assert 500 * ((1 + i) ** n - 1) / i == pytest.approx(60_000)
    assert co._anios_para(1_000, 60_000, 0.07) < anios


def test_cada_pregunta_apunta_a_un_documento_del_compilado():
    codigos = {d[0] for d in co.DOCUMENTOS}
    assert {p[3] for p in co.PREGUNTAS} <= codigos
    assert len(co.CONCLUSIONES) == 10
    assert all(t in {"favorable", "desfavorable", "neutral"} for _, _, t in co.CONCLUSIONES)


def test_las_barras_dejan_espacio_para_la_cifra():
    fig, _ = co._barras_h(["a", "b"], [-0.049, 0.0092], titulo="t")
    lo, hi = fig.layout.xaxis.range
    assert lo < -0.049 and hi > 0.0092


def test_el_plan_del_excel_es_formula_viva(tmp_path):
    ruta = xc.exportar(tmp_path / "compilado.xlsx", resultados.cargar())
    wb = openpyxl.load_workbook(ruta)
    assert wb.calculation.fullCalcOnLoad
    ws = wb["Mi plan"]
    # Los insumos van en azul sobre amarillo y los resultados son fórmulas que los leen.
    for i in range(len(xc.INSUMOS_PLAN)):
        c = ws.cell(row=xc.FILA_INSUMOS + i, column=2)
        assert c.fill.fgColor.rgb == RELLENO_INPUT.fgColor.rgb and c.value is not None
    inicio = xc.FILA_INSUMOS + len(xc.INSUMOS_PLAN) + 1
    formulas = [ws.cell(row=inicio + 1 + i, column=2).value for i in range(len(xc.RESULTADOS_PLAN))]
    assert all(isinstance(f, str) and f.startswith("=") for f in formulas)
    assert formulas[0] == f"={xc.CELDA['gastos']}*{xc.CELDA['meses']}"
    assert "{" not in "".join(formulas)                           # no quedó ningún hueco sin llenar
    assert ws.cell(row=xc.FILA_INSUMOS + 7, column=2).value == xc.precio_schh() > 0
    assert {"Preguntas", "Conclusiones", "Mi plan", "Cuándo entrar", "En cuáles", "Retorno esperado", "ETFs",
            "Satélite", "Edges", "Impuestos", "Glosario"} == set(wb.sheetnames)


def test_el_retorno_esperado_sube_con_lo_barato():
    medianas = [r[2] for r in co.RETORNO_ESPERADO]
    assert medianas == sorted(medianas)
    assert all(lo < m < hi for _, _, m, lo, hi in co.RETORNO_ESPERADO)
    assert np.all(np.array(medianas) > co.EFECTIVO_5A)
