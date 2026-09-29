"""Modelos de valor intrínseco en el tiempo (``src/estudio/intrinsecos.py``).

Primero las fórmulas, contra cuentas hechas a mano: un DCF de dos etapas con el mismo
crecimiento en las dos es exactamente un Gordon, y un Gordon valuado con el crecimiento
implícito devuelve el precio. Después los datos del NAV, que salen de XBRL versionado y
tienen que respetar la fecha de publicación de cada cifra (P1).
"""

from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

RAIZ = Path(__file__).resolve().parent.parent
for ruta in (str(RAIZ), str(RAIZ / "app")):
    if ruta not in sys.path:
        sys.path.insert(0, ruta)

from src.estudio import intrinsecos as it  # noqa: E402


def test_el_ddm_es_gordon():
    assert it.ddm(1.0, 0.03, 0.08) == pytest.approx(1.03 / 0.05)
    assert np.isnan(it.ddm(1.0, 0.08, 0.08)), "con r = g no hay valor, hay división entre cero"
    assert np.isnan(it.ddm(0.0, 0.02, 0.08))


def test_dos_etapas_con_el_mismo_crecimiento_es_gordon():
    assert it.dcf_dos_etapas(2.0, 0.02, 0.02, 0.08, 5) == pytest.approx(it.ddm(2.0, 0.02, 0.08), rel=1e-12)


def test_dos_etapas_a_mano():
    r, g1, g2 = 0.08, 0.05, 0.02
    flujos = [1.05, 1.05**2, 1.05**3, 1.05**4, 1.05**5]
    esperado = sum(f / 1.08**k for k, f in enumerate(flujos, start=1)) + flujos[-1] * 1.02 / 0.06 / 1.08**5
    assert it.dcf_dos_etapas(1.0, g1, g2, r, 5) == pytest.approx(esperado, rel=1e-12)


def test_el_crecimiento_implicito_devuelve_el_precio():
    precio, flujo, r = 55.0, 4.2, 0.0817
    g = it.crecimiento_implicito_seguro(precio, flujo, r)
    assert flujo * (1 + g) / (r - g) == pytest.approx(precio, rel=1e-12)


def test_acotar_respeta_limites_y_margen():
    assert it.acotar(0.07, (0.0, 0.04)) == 0.04
    assert it.acotar(-0.02, (0.0, 0.04)) == 0.0
    assert it.acotar(0.04, (0.0, 0.04), r=0.05, margen=0.02) == pytest.approx(0.03)
    assert np.isnan(it.acotar(np.nan, (0.0, 0.04)))


def test_crecimiento_anual_compuesto():
    fechas = pd.date_range("2010-01-31", "2020-12-31", freq="ME")
    serie = pd.Series(2.0 ** ((fechas - fechas[0]).days / 365.25 / 5), index=fechas)
    g = it.crecimiento_anual(serie, fechas, 5)
    assert g.iloc[:59].isna().all(), "sin 5 años de historia no hay crecimiento"
    assert g.iloc[-1] == pytest.approx(2 ** (1 / 5) - 1, abs=1e-3)


def test_la_senal_absoluta_usa_el_margen():
    assert it.senal_absoluta(0.15, 0.15) == "barato"
    assert it.senal_absoluta(0.1499, 0.15) == "medio"
    assert it.senal_absoluta(-0.15, 0.15) == "caro"
    assert it.senal_absoluta(np.nan, 0.15) == "sin dato"


def test_el_diseno_quedo_fijado():
    """Los supuestos escritos antes de correr. Si esta prueba cambia, el cambio va en el
    historial con su motivo: moverlos después de ver resultados sería ajustar el modelo
    al pasado."""
    s = it.SUPUESTOS
    assert (s.prima, s.anios_historia, s.g_dividendo, s.g_flujo, s.g_terminal, s.anios_etapa_1,
            s.margen_r_menos_g, s.g_entregado, s.umbral_valor, s.umbral_crecimiento) == (
        0.03, 5, (0.0, 0.04), (0.0, 0.08), 0.02, 5, 0.02, (-0.05, 0.10), 0.15, 0.01)
    assert it.CLAVES == ("ddm", "dcf", "crecimiento", "nav")


# --------------------------------------------------------------------------------------
# El NAV desde XBRL versionado
# --------------------------------------------------------------------------------------

import datetime as dt  # noqa: E402

from src.datos import almacen  # noqa: E402
from src.estudio import mercado  # noqa: E402

D = dt.date


def _crudos(ticker: str) -> pd.DataFrame:
    try:
        c = almacen.leer_crudos(ticker)
    except (FileNotFoundError, OSError):
        pytest.skip(f"No están versionados los hechos XBRL de {ticker}.")
    if c.empty:
        pytest.skip(f"No están versionados los hechos XBRL de {ticker}.")
    return c


def test_la_deuda_suma_sus_componentes():
    """XBRL etiqueta como «total» solo los bonos de Realty Income (25,092 millones); con
    el préstamo, la línea y el papel comercial son 30,652."""
    n = it.nav_conocido(_crudos("O"), "O", pd.DatetimeIndex(["2026-09-28"]))
    assert n["trimestre"].iloc[0] == pd.Timestamp("2026-06-30")
    assert n["deuda"].iloc[0] / 1e6 == pytest.approx(30_651.7, abs=1.0)


def test_el_total_etiquetado_no_manda_si_la_suma_es_mayor():
    """A fines de 2016 Realty Income etiquetó como deuda de largo plazo solo sus bonos (3,975
    millones); con la línea de crédito, los préstamos y las otras notas eran 5,880."""
    n = it.nav_conocido(_crudos("O"), "O", pd.DatetimeIndex(["2017-03-31"]))
    assert n["trimestre"].iloc[0] == pd.Timestamp("2016-12-31")
    assert n["deuda"].iloc[0] / 1e6 == pytest.approx(5_880.2, abs=0.5)


def test_la_etiqueta_vieja_de_la_linea_de_credito_cuenta():
    """En 2012 NNN reportaba su línea de crédito como ``LineOfCredit`` (142.6 millones); hoy usa
    otra etiqueta. Elegir la de hoy para toda la historia la dejaba fuera."""
    n = it.nav_conocido(_crudos("NNN"), "NNN", pd.DatetimeIndex(["2012-09-28"]))
    assert n["trimestre"].iloc[0] == pd.Timestamp("2012-06-30")
    assert n["deuda"].iloc[0] / 1e6 == pytest.approx(29.3 + 845.2 + 142.6, abs=0.5)


def test_un_hueco_de_etiquetas_queda_vacio():
    """O etiquetó sus bonos de 2017 a mediados de 2018 con una etiqueta propia: la deuda armada
    no llega al 60% del pasivo y el mes queda vacío en vez de salir con un NAV inflado."""
    n = it.nav_conocido(_crudos("O"), "O", pd.DatetimeIndex(["2018-09-28"]))
    assert np.isnan(n["noi"].iloc[0])


def test_noi_y_balance_salen_del_mismo_trimestre():
    """Realty Income compró VEREIT en noviembre de 2021. A fines de diciembre lo último publicado
    era el tercer trimestre: NOI y deuda de ANTES de la compra, los dos. En marzo de 2022, con el
    10-K, el cuarto trimestre: los dos ya con VEREIT. Mezclarlos —deuda nueva con NOI viejo— dejaba
    el NAV a la mitad."""
    n = it.nav_conocido(_crudos("O"), "O", pd.DatetimeIndex(["2021-12-31", "2022-03-31"]))
    antes, despues = n.iloc[0], n.iloc[1]
    assert antes["trimestre"] == pd.Timestamp("2021-09-30") and antes["deuda"] < 10e9
    assert despues["trimestre"] == pd.Timestamp("2021-12-31") and despues["deuda"] > 13e9
    assert despues["noi"] * 4 > 2.5e9, "el NOI del mismo corte ya trae los inmuebles de VEREIT"


def test_el_nav_no_mira_al_futuro():
    """Con los hechos publicados hasta 2019, el NAV de 2012–2019 es el mismo que con todo (P1)."""
    c = _crudos("NNN")
    fechas = pd.date_range("2012-01-31", "2019-12-31", freq="ME")
    completo = it.nav_conocido(c, "NNN", fechas)
    hasta = c[pd.to_datetime(c["fecha_publicacion"]) <= "2019-12-31"]
    cortado = it.nav_conocido(hasta, "NNN", fechas)
    pd.testing.assert_frame_equal(completo, cortado)


def _armar(repo, ticker: str, asof: dt.date):
    if not mercado.hay_estudio(ticker):
        pytest.skip(f"No está versionado el estudio de {ticker}.")
    from src.estudio import estudio

    return estudio.armar(repo, ticker, asof=asof)


def test_el_apalancamiento_armado_cuadra_con_el_reportado(repo_sembrado):
    """La validación que decide si el NAV sirve: NNN desde 2020 y WPC desde 2019 a menos de 12% de
    lo que reportó cada emisor; WPC antes de 2019 no se usa."""
    for ticker in ("NNN", "WPC"):
        e = _armar(repo_sembrado, ticker, D(2026, 9, 28))
        fechas = pd.date_range("2013-01-31", "2026-08-31", freq="ME")
        d = it.insumos_nav(e, fechas)
        v = it.validacion_nav(e, d)
        recientes = v[v["anio"] >= 2020]
        assert not recientes.empty
        assert (recientes["diferencia"].abs() < 0.12).all(), recientes
        if ticker == "WPC":
            assert d.loc[d["trimestre_nav"] < "2019-03-31", "nav"].isna().all()
            assert not v.loc[v["anio"] < 2019, "se_usa"].any()


# --------------------------------------------------------------------------------------
# Los cuatro modelos sobre la historia versionada
# --------------------------------------------------------------------------------------


def test_los_modelos_no_miran_al_futuro(repo_sembrado):
    from src.estudio import metodos

    tarde = _armar(repo_sembrado, "O", D(2026, 6, 30))
    temprano = _armar(repo_sembrado, "O", D(2015, 6, 30))
    a = it.agregar(metodos.panel(tarde), tarde)
    b = it.agregar(metodos.panel(temprano), temprano)
    columnas = [f"{x}_{c}" for c in it.CLAVES for x in ("v", "p", "a")] + ["valor_ddm", "valor_dcf", "nav"]
    pd.testing.assert_frame_equal(a.loc[:"2015-05-31", columnas], b.loc[:"2015-05-31", columnas])


@pytest.fixture(scope="module")
def r_intrinsecos(tmp_path_factory):
    from src.datos.repositorio import Repositorio
    from src.datos.semilla import sembrar
    from src.estudio import metodos

    repo = Repositorio(ruta=tmp_path_factory.mktemp("base") / "sembrada.db")
    sembrar(repo, tickers=["O", "PLD", "GNL"], inicio=D(2018, 1, 1), fin=D(2026, 6, 30))
    tickers = [t for t in ("O", "NNN") if mercado.hay_estudio(t)]
    if len(tickers) < 2:
        pytest.skip("Hacen falta dos emisores con estudio versionado.")
    return metodos.estudiar({t: _armar(repo, t, D(2026, 6, 30)) for t in tickers}, n_azar=5)


def test_la_senal_absoluta_y_el_percentil_salen_de_la_misma_serie(r_intrinsecos):
    ri = r_intrinsecos.intrinsecos
    for p in ri.paneles.values():
        con_valor = p["v_ddm"].notna()
        assert (p.loc[~con_valor, "a_ddm"] == "sin dato").all()
        barato = p["v_ddm"] >= it.SUPUESTOS.umbral_valor
        assert (p.loc[barato, "a_ddm"] == "barato").all()
        # Con la prima de 3% la sensibilidad es la señal base.
        assert (p["a_ddm_prima_300"] == p["a_ddm"]).all()


def test_las_conclusiones_de_los_modelos_estan_completas(r_intrinsecos):
    import re

    c = it.conclusiones(r_intrinsecos.intrinsecos, r_intrinsecos)
    assert len(c) == 5
    assert c[-1].titulo == "Lo que se escribió antes de correr"
    for x in c:
        assert not re.search(r"\bnan\b|—(?: pb|%| al año)", x.texto), x.texto


def test_las_vistas_de_los_modelos_ponen_cada_fraccion_en_porcentaje(r_intrinsecos):
    from comun import familia_de_columna

    from src.estudio import vistas_intrinsecos as vi

    ri = r_intrinsecos.intrinsecos
    t = next(iter(ri.paneles))
    tablas = {"modelos": vi.modelos(), "hoy": vi.hoy(ri), "evaluacion": vi.evaluacion(ri),
              "sensibilidad": vi.sensibilidad(ri), "asignacion": vi.asignacion(ri),
              "validacion": vi.validacion(ri), "backtest": vi.backtest(ri, t), "trimestres": vi.trimestres(ri, t)}
    problemas = []
    for nombre, tabla in tablas.items():
        if list(tabla.columns) != list(vi.COLUMNAS[nombre]):
            problemas.append(f"{nombre}: columnas {list(tabla.columns)}")
        for columna in tabla.columns:
            serie = pd.to_numeric(tabla[columna], errors="coerce").dropna()
            if serie.empty or tabla[columna].dtype == object or pd.api.types.is_integer_dtype(tabla[columna]):
                continue
            if columna.startswith("correlacion"):
                continue
            if serie.abs().max() < 1 and familia_de_columna(columna, serie) != "porcentaje":
                problemas.append(f"{nombre}.{columna}: fracción en «{familia_de_columna(columna, serie)}»")
    assert not problemas, "\n".join(problemas)


_SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


@pytest.mark.skipif(_SOFFICE is None, reason="LibreOffice no está instalado.")
def test_el_excel_reproduce_los_cuatro_modelos(r_intrinsecos, tmp_path):
    """Cada renglón de la hoja «O valor», recalculado por LibreOffice, da lo mismo que el estudio."""
    from openpyxl import load_workbook

    from src.estudio.metodos import trimestral
    from src.export import excel_metodos
    from src.export.excel import buscar_errores

    ruta = excel_metodos.exportar(r_intrinsecos, tmp_path / "valuacion_trimestral.xlsx")
    salida = tmp_path / "recalc"
    subprocess.run([_SOFFICE, "--headless", f"-env:UserInstallation=file://{tmp_path}/perfil",
                    "--convert-to", "xlsx", "--outdir", str(salida), str(ruta)],
                   capture_output=True, timeout=300, check=False)
    recalculado = salida / ruta.name
    assert buscar_errores(recalculado) == []
    ws = load_workbook(recalculado, data_only=True)["O valor"]
    encabezado = [c.value for c in ws[1]]
    q = trimestral(r_intrinsecos.intrinsecos.paneles["O"]).dropna(subset=["precio"])
    q = q[q["r"].notna()]
    for fila, (_, esperado) in zip(ws.iter_rows(min_row=2), q.iterrows(), strict=True):
        x = {h: c.value for h, c in zip(encabezado, fila, strict=True)}
        for h, col in (("Valor DDM", "valor_ddm"), ("Valor DCF", "valor_dcf"), ("NAV por acción", "nav"),
                       ("Entregado − implícito", "v_crecimiento")):
            if pd.isna(esperado[col]):
                assert x[h] in (None, ""), (h, x[h])
            else:
                assert x[h] == pytest.approx(esperado[col], rel=1e-9, abs=1e-12), h
        for h, col in (("Señal DDM", "a_ddm"), ("Señal DCF", "a_dcf"), ("Señal crecimiento", "a_crecimiento"),
                       ("Señal NAV", "a_nav")):
            assert x[h] == esperado[col], h
